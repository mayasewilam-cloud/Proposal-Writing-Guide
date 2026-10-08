"""
Automated verification suite for the Extended Faculty Portal & Shared Team Workspaces:
- Attendance CSV & XLSX Upload, Preview (Duplicates/Missing/New/Existing), & Confirmation
- Student Directory & Multi-Criteria Filtering (Lab Group, Team, PI, Assignment Status, Search)
- Lab Group Management & Student Lab Group Move (L1 -> L3) with Immutable Audit History
- Project Team Setup inside Lab Groups, Unassigned Students Tracking, & Single-Active-Team Rule
- PI / Supervisor Assignment & Official Topic Assignment
- Shared Team Workspace persisted by Team ID (Teammate A writes -> Teammate B, PI & Faculty see it; Other Team blocked with 403)
- Strict Role-Based Access Control (Students blocked from administrative mutations)
"""

import io
import json
import os
import shutil
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

import app as portal_app_module
from core import auth_store, cohort_store


class FacultyCohortAndSharedTeamWorkspaceTestCase(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.mkdtemp(prefix="biowriter_cohort_test_")
        self.orig_cohort_db_file = cohort_store.COHORT_DB_FILE
        self.orig_team_ws_dir = cohort_store.TEAM_WORKSPACES_DIR
        self.orig_users_file = auth_store.USERS_FILE
        self.orig_cohort_users_file = cohort_store.USERS_FILE

        cohort_store.COHORT_DB_FILE = os.path.join(self.temp_dir, "cohort_db.json")
        cohort_store.TEAM_WORKSPACES_DIR = os.path.join(self.temp_dir, "team_workspaces")
        temp_users_file = os.path.join(self.temp_dir, "users.json")
        with open(temp_users_file, "w", encoding="utf-8") as f:
            json.dump({}, f)
        auth_store.USERS_FILE = temp_users_file
        cohort_store.USERS_FILE = temp_users_file

        portal_app_module.app.config["TESTING"] = True
        self.client = portal_app_module.app.test_client()

    def tearDown(self):
        cohort_store.COHORT_DB_FILE = self.orig_cohort_db_file
        cohort_store.TEAM_WORKSPACES_DIR = self.orig_team_ws_dir
        auth_store.USERS_FILE = self.orig_users_file
        cohort_store.USERS_FILE = self.orig_cohort_users_file
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def _login_faculty(self):
        resp = self.client.post(
            "/api/faculty/login",
            json={"passcode": "faculty2026"},
        )
        self.assertEqual(resp.status_code, 200)
        self.assertTrue(resp.get_json()["success"])

    def _login_student(self, student_id, password="password123"):
        resp = self.client.post(
            "/api/auth/login",
            json={"identifier": student_id, "password": password},
        )
        self.assertEqual(resp.status_code, 200)
        data = resp.get_json()
        self.assertTrue(data["success"])
        return data

    def _login_pi(self, pi_id="PI-RANIA", passcode="pi2026"):
        resp = self.client.post(
            "/api/pi/login",
            json={"pi_id": pi_id, "passcode": passcode},
        )
        self.assertEqual(resp.status_code, 200)
        data = resp.get_json()
        self.assertTrue(data["success"])
        return data

    # =========================================================================
    # 1. ATTENDANCE CSV & XLSX PREVIEW (WITHOUT OVERWRITE) + CONFIRMATION
    # =========================================================================
    def test_01_attendance_upload_preview_and_confirm_csv_and_xlsx(self):
        self._login_faculty()

        # Check initial student count before preview
        before_resp = self.client.get("/api/faculty/students")
        before_count = before_resp.get_json()["total"]

        # Upload CSV containing New, Existing, Duplicate Row, Duplicate ID, and Missing Fields
        csv_content = "\n".join([
            "Student ID,Student Name,Lab Group,Email",
            "20230001,Ahmed Ali,L1,ahmed.ali@student.edu.eg",            # Existing
            "20230901,Nadia Mansour,L1,nadia.mansour@student.edu.eg",    # New
            "20230902,Tamer Hosny,L4,tamer.hosny@student.edu.eg",        # New + new lab group L4
            "20230901,Nadia Mansour,L1,nadia.mansour@student.edu.eg",    # Duplicate Row
            "20230902,Different Name Same ID,L2,dup@student.edu.eg",     # Duplicate Student ID
            ",Missing ID Student,L1,noid@student.edu.eg",                # Missing Student ID
            "20230903,,L1,noname@student.edu.eg",                        # Missing Student Name
            "20230904,Salma Without Lab,,nolab@student.edu.eg",          # Missing Lab Group
        ])

        preview_resp = self.client.post(
            "/api/faculty/attendance/preview",
            data={
                "semester": "Fall 2026",
                "course_id": "BT_301",
                "file": (io.BytesIO(csv_content.encode("utf-8")), "cohort_attendance.csv"),
            },
            content_type="multipart/form-data",
        )
        self.assertEqual(preview_resp.status_code, 200)
        prev = preview_resp.get_json()
        self.assertTrue(prev["success"])
        self.assertEqual(prev["students_detected"], 3)  # 20230001, 20230901, 20230902
        self.assertEqual(prev["lab_groups_detected"], 2)  # L1, L4
        self.assertEqual(prev["summary"]["existing_count"], 1)
        self.assertEqual(prev["summary"]["new_count"], 2)
        self.assertEqual(prev["summary"]["duplicate_row_count"], 1)
        self.assertEqual(prev["summary"]["duplicate_id_count"], 1)
        self.assertEqual(prev["summary"]["missing_id_count"], 1)
        self.assertEqual(prev["summary"]["missing_name_count"], 1)
        self.assertEqual(prev["summary"]["missing_lab_group_count"], 1)

        # Ensure data was NOT overwritten before confirmation
        mid_resp = self.client.get("/api/faculty/students")
        self.assertEqual(mid_resp.get_json()["total"], before_count)

        # Confirm the attendance import
        confirm_resp = self.client.post(
            "/api/faculty/attendance/confirm",
            json={"preview_id": prev["preview_id"]},
        )
        self.assertEqual(confirm_resp.status_code, 200)
        conf = confirm_resp.get_json()
        self.assertTrue(conf["success"])
        self.assertEqual(conf["imported_count"], 3)
        self.assertEqual(conf["created_count"], 2)
        self.assertEqual(conf["updated_count"], 1)

        # Verify database now includes the 2 new students
        after_resp = self.client.get("/api/faculty/students")
        self.assertEqual(after_resp.get_json()["total"], before_count + 2)

        # Also test .xlsx file upload parsing
        try:
            import openpyxl
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(["Student ID", "Student Name", "Lab Group"])
            ws.append(["20230999", "Excel Imported Student", "L3"])
            xlsx_buf = io.BytesIO()
            wb.save(xlsx_buf)
            xlsx_buf.seek(0)

            xlsx_prev = self.client.post(
                "/api/faculty/attendance/preview",
                data={"file": (xlsx_buf, "sheet.xlsx")},
                content_type="multipart/form-data",
            )
            self.assertEqual(xlsx_prev.status_code, 200)
            x_data = xlsx_prev.get_json()
            self.assertTrue(x_data["success"])
            self.assertEqual(x_data["students_detected"], 1)
            self.assertEqual(x_data["rows"][0]["student_id"], "20230999")
            self.assertEqual(x_data["rows"][0]["lab_group"], "L3")
        except ImportError:
            pass

    # =========================================================================
    # 2. LAB GROUP MOVE (L1 -> L3) & AUDIT HISTORY LOGGING
    # =========================================================================
    def test_02_move_student_lab_group_and_audit_history(self):
        self._login_faculty()

        # Ahmed Ali (20230001) starts in L1 (Team L1-T1)
        move_resp = self.client.post(
            "/api/faculty/students/move_lab_group",
            json={
                "student_id": "20230001",
                "new_lab_group": "L3",
            },
        )
        self.assertEqual(move_resp.status_code, 200)
        mv = move_resp.get_json()
        self.assertTrue(mv["success"])
        self.assertEqual(mv["old_lab_group"], "L1")
        self.assertEqual(mv["new_lab_group"], "L3")
        self.assertIn("Moved from L1 to L3", mv["audit_entry"]["summary"])

        # Verify student record in database is updated to L3 and detached from L1-T1
        st_resp = self.client.get("/api/faculty/students?q=20230001")
        st_list = st_resp.get_json()["students"]
        self.assertEqual(len(st_list), 1)
        self.assertEqual(st_list[0]["lab_group"], "L3")
        self.assertEqual(st_list[0]["team_id"], "")

        # Verify audit history contains the LAB_GROUP_MOVE event
        hist_resp = self.client.get("/api/faculty/audit_history?student_id=20230001")
        hist = hist_resp.get_json()["history"]
        self.assertTrue(any(h["event_type"] == "LAB_GROUP_MOVE" and h["new_value"] == "L3" for h in hist))

    # =========================================================================
    # 3. PROJECT TEAMS SETUP, UNASSIGNED STUDENTS & SINGLE-ACTIVE-TEAM RULE
    # =========================================================================
    def test_03_team_creation_unassigned_tracking_and_single_active_team_rule(self):
        self._login_faculty()

        # Check unassigned students in L1 (20230005 Omar Khaled & 20230006 Nour Hany are unassigned in L1)
        lg_resp = self.client.get("/api/faculty/lab_groups")
        l1 = next(g for g in lg_resp.get_json()["lab_groups"] if g["lab_group_id"] == "L1")
        unassigned_ids = [u["student_id"] for u in l1["unassigned_students"]]
        self.assertIn("20230005", unassigned_ids)
        self.assertIn("20230006", unassigned_ids)

        # Create a new project team in L1 with those unassigned students
        create_resp = self.client.post(
            "/api/faculty/teams",
            json={
                "lab_group": "L1",
                "team_name": "Team 3 — Biosensor Delta",
                "student_ids": ["20230005", "20230006"],
                "pi_id": "PI-AHMED",
                "topic_title": "CRISPR Phage Display Diagnostics",
            },
        )
        self.assertEqual(create_resp.status_code, 200)
        t_data = create_resp.get_json()
        self.assertTrue(t_data["success"])
        new_team_id = t_data["team"]["team_id"]
        self.assertEqual(t_data["team"]["member_count"], 2)

        # Verify 20230005 is no longer in L1 unassigned_students
        lg_after = self.client.get("/api/faculty/lab_groups").get_json()["lab_groups"]
        l1_after = next(g for g in lg_after if g["lab_group_id"] == "L1")
        self.assertNotIn("20230005", [u["student_id"] for u in l1_after["unassigned_students"]])

        # SINGLE-ACTIVE-TEAM ENFORCEMENT:
        # Attempt to add 20230001 (already in active team L1-T1) to new_team_id without allow_move
        dup_team_resp = self.client.post(
            f"/api/faculty/teams/{new_team_id}/members",
            json={"student_id": "20230001", "allow_move": False},
        )
        self.assertEqual(dup_team_resp.status_code, 400)
        self.assertIn("only one active project team", dup_team_resp.get_json()["error"].lower())

        # Attempt to create another team containing 20230001 (already in L1-T1) -> must be rejected
        bad_create = self.client.post(
            "/api/faculty/teams",
            json={
                "lab_group": "L1",
                "team_name": "Invalid Duplicate Member Team",
                "student_ids": ["20230001"],
            },
        )
        self.assertEqual(bad_create.status_code, 400)
        self.assertIn("only one active project team", bad_create.get_json()["error"].lower())

    # =========================================================================
    # 4. PI ASSIGNMENT & TOPIC ASSIGNMENT
    # =========================================================================
    def test_04_assign_pi_and_topic_to_team(self):
        self._login_faculty()

        # Assign PI-MONA to L2-T1
        pi_resp = self.client.post(
            "/api/faculty/teams/L2-T1/assign_pi",
            json={"pi_id": "PI-MONA"},
        )
        self.assertEqual(pi_resp.status_code, 200)
        self.assertEqual(pi_resp.get_json()["team"]["pi_id"], "PI-MONA")
        self.assertEqual(pi_resp.get_json()["team"]["pi_name"], "Dr. Mona Abdelrahman")

        # Assign official Topic to L2-T1
        topic_resp = self.client.post(
            "/api/faculty/teams/L2-T1/assign_topic",
            json={
                "topic_title": "Synthetic Quorum Quenching Lactonases for Biofilm Inhibition",
                "topic_description": "Engineered AiiA lactonase variants against P. aeruginosa",
                "topic_category": "Synthetic Biology",
            },
        )
        self.assertEqual(topic_resp.status_code, 200)
        self.assertEqual(
            topic_resp.get_json()["team"]["topic_title"],
            "Synthetic Quorum Quenching Lactonases for Biofilm Inhibition",
        )

        # Verify Shared Workspace for L2-T1 reflects the new PI and Topic immediately
        ws_resp = self.client.get("/api/teams/L2-T1/workspace")
        ws = ws_resp.get_json()["workspace"]
        self.assertEqual(ws["pi_id"], "PI-MONA")
        self.assertEqual(ws["topic_title"], "Synthetic Quorum Quenching Lactonases for Biofilm Inhibition")

    # =========================================================================
    # 5. SHARED TEAM WORKSPACE SYNCHRONIZATION & CROSS-TEAM ISOLATION
    # =========================================================================
    def test_05_shared_team_workspace_collaboration_and_access_control(self):
        # Student A (20230001, Ahmed Ali) logs in -> belongs to team L1-T1
        login_a = self._login_student("20230001")
        self.assertEqual(login_a["my_project"]["team_id"], "L1-T1")
        self.assertEqual(login_a["my_project"]["lab_group"], "L1")

        # Student A writes shared notes, important info, and an update to L1-T1 workspace
        save_resp = self.client.post(
            "/api/teams/L1-T1/workspace",
            json={
                "shared_notes": "Aim 1 crRNA sequences verified by Ahmed Ali. Next: Sara runs LoD probit.",
                "important_info": "Biosafety Level 2 cabinet reserved Mondays & Wednesdays.",
                "update_message": "Ahmed added Aim 1 crRNA validation notes",
                "proposal_draft": {
                    "title": "DETECTR-Aqua Team Updated Proposal Title",
                    "abstract": "Shared team abstract updated by Student A for all L1-T1 teammates.",
                },
            },
        )
        self.assertEqual(save_resp.status_code, 200)
        self.assertTrue(save_resp.get_json()["success"])

        # Student B (20230002, Sara Mohamed, ALSO in team L1-T1) logs in and fetches workspace
        login_b = self._login_student("20230002")
        self.assertEqual(login_b["my_project"]["team_id"], "L1-T1")

        ws_b_resp = self.client.get("/api/teams/L1-T1/workspace")
        self.assertEqual(ws_b_resp.status_code, 200)
        ws_b = ws_b_resp.get_json()["workspace"]
        self.assertEqual(
            ws_b["shared_notes"],
            "Aim 1 crRNA sequences verified by Ahmed Ali. Next: Sara runs LoD probit.",
        )
        self.assertEqual(
            ws_b["important_info"],
            "Biosafety Level 2 cabinet reserved Mondays & Wednesdays.",
        )
        self.assertEqual(
            ws_b["proposal_draft"]["title"],
            "DETECTR-Aqua Team Updated Proposal Title",
        )
        self.assertTrue(
            any("Ahmed added Aim 1 crRNA validation notes" in u["message"] for u in ws_b["updates_feed"])
        )

        # Assigned PI (PI-RANIA) logs in and views supervised teams + L1-T1 shared workspace
        self._login_pi("PI-RANIA", "pi2026")
        pi_teams_resp = self.client.get("/api/pi/my_teams")
        self.assertEqual(pi_teams_resp.status_code, 200)
        pi_team_ids = [t["team_id"] for t in pi_teams_resp.get_json()["teams"]]
        self.assertIn("L1-T1", pi_team_ids)

        ws_pi_resp = self.client.get("/api/teams/L1-T1/workspace")
        self.assertEqual(ws_pi_resp.status_code, 200)
        self.assertEqual(
            ws_pi_resp.get_json()["workspace"]["shared_notes"],
            "Aim 1 crRNA sequences verified by Ahmed Ali. Next: Sara runs LoD probit.",
        )

        # Student C (20230007, Youssef Tarek, in L2-T1) logs in and attempts to access L1-T1 workspace -> 403 Forbidden
        self._login_student("20230007")
        forbidden_get = self.client.get("/api/teams/L1-T1/workspace")
        self.assertEqual(forbidden_get.status_code, 403)

        forbidden_post = self.client.post(
            "/api/teams/L1-T1/workspace",
            json={"shared_notes": "Unauthorized cross-team write attempt"},
        )
        self.assertEqual(forbidden_post.status_code, 403)

    # =========================================================================
    # 6. PERMISSION ENFORCEMENT: STUDENTS CANNOT PERFORM FACULTY ACTIONS
    # =========================================================================
    def test_06_student_cannot_modify_official_assignments(self):
        self._login_student("20230001")

        # Student cannot upload attendance
        r1 = self.client.post("/api/faculty/attendance/preview", json={"csv_text": "Student ID,Name,Lab\n1,A,L1"})
        self.assertEqual(r1.status_code, 403)

        # Student cannot move lab group
        r2 = self.client.post(
            "/api/faculty/students/move_lab_group",
            json={"student_id": "20230001", "new_lab_group": "L3"},
        )
        self.assertEqual(r2.status_code, 403)

        # Student cannot create or modify teams
        r3 = self.client.post(
            "/api/faculty/teams",
            json={"lab_group": "L1", "team_name": "Unauthorized Team"},
        )
        self.assertEqual(r3.status_code, 403)

        # Student cannot change PI or Topic
        r4 = self.client.post("/api/faculty/teams/L1-T1/assign_pi", json={"pi_id": "PI-AHMED"})
        self.assertEqual(r4.status_code, 403)

        r5 = self.client.post("/api/faculty/teams/L1-T1/assign_topic", json={"topic_title": "Unauthorized Topic"})
        self.assertEqual(r5.status_code, 403)

        # Student cannot inject official assignment fields via team workspace endpoint
        r6 = self.client.post(
            "/api/teams/L1-T1/workspace",
            json={"pi_id": "PI-AHMED", "shared_notes": "Trying to hijack PI"},
        )
        self.assertEqual(r6.status_code, 403)

    # =========================================================================
    # 7. PRE-ASSIGNED STUDENTS CREATE ACCOUNT WITH THEIR ID & SEE TEAM + MEMBERS
    # =========================================================================
    def test_07_pre_enrolled_student_creates_account_and_sees_assigned_team_and_members(self):
        self._login_faculty()

        # 1. Faculty uploads attendance sheet with 3 new students in Lab Group L4
        csv_data = "\n".join([
            "Student ID,Student Name,Lab Group,Email",
            "20230901,Hana Tarek,L4,hana.tarek@student.edu.eg",
            "20230902,Ziad Adel,L4,ziad.adel@student.edu.eg",
            "20230903,Nouran Sameh,L4,nouran.sameh@student.edu.eg",
        ])
        prev_resp = self.client.post(
            "/api/faculty/attendance/preview",
            json={"csv_text": csv_data, "semester": "Fall 2026", "course_id": "BT_301"},
        )
        self.assertEqual(prev_resp.status_code, 200)
        preview_id = prev_resp.get_json()["preview_id"]

        conf_resp = self.client.post(
            "/api/faculty/attendance/confirm",
            json={"preview_id": preview_id},
        )
        self.assertEqual(conf_resp.status_code, 200)
        self.assertTrue(conf_resp.get_json()["success"])

        # 2. Faculty selects Lab Group L4, selects the 3 students, states they are Team 1, and writes their Topic
        create_team_resp = self.client.post(
            "/api/faculty/teams",
            json={
                "lab_group": "L4",
                "team_name": "Team 1 — AptamerSense",
                "student_ids": ["20230901", "20230902", "20230903"],
                "pi_id": "PI-RANIA",
                "topic_title": "Electrochemical Aptamer Biosensor for Aflatoxin M1 in Dairy",
                "topic_description": "Screen-printed carbon electrode functionalized with DNA aptamer.",
                "semester": "Fall 2026",
                "course_id": "BT_301",
            },
        )
        self.assertEqual(create_team_resp.status_code, 200)
        team_data = create_team_resp.get_json()["team"]
        self.assertEqual(team_data["team_id"], "L4-T1")
        self.assertEqual(team_data["topic_title"], "Electrochemical Aptamer Biosensor for Aflatoxin M1 in Dairy")

        # 3. Sign out faculty, then student 20230901 goes to Student Portal -> Create Account (/api/auth/register)
        self.client.post("/api/faculty/logout")
        reg_resp = self.client.post(
            "/api/auth/register",
            json={
                "student_id": "20230901",
                "student_name": "Hana Tarek",
                "password": "hanaSecretPassword99",
            },
        )
        self.assertEqual(reg_resp.status_code, 200)
        reg_json = reg_resp.get_json()
        self.assertTrue(reg_json["success"])
        self.assertEqual(reg_json["user"]["student_id"], "20230901")

        # Verify student immediately sees their Lab Group, Team, Topic, PI, and all Team Members
        my_proj = reg_json["my_project"]
        self.assertIsNotNone(my_proj)
        self.assertEqual(my_proj["lab_group"], "L4")
        self.assertEqual(my_proj["team_id"], "L4-T1")
        self.assertEqual(my_proj["team_name"], "Team 1 — AptamerSense")
        self.assertEqual(my_proj["pi_name"], "Dr. Rania Hassan")
        self.assertEqual(my_proj["topic_title"], "Electrochemical Aptamer Biosensor for Aflatoxin M1 in Dairy")
        member_ids = [m["student_id"] for m in my_proj["members"]]
        self.assertEqual(member_ids, ["20230901", "20230902", "20230903"])

        # Verify shared draft also contains the team roster and topic
        self.assertEqual(reg_json["draft"]["title"], "Electrochemical Aptamer Biosensor for Aflatoxin M1 in Dairy")
        self.assertEqual(len(reg_json["draft"]["team_roster"]), 3)

        # Verify trying to register the same claimed account a second time is rejected
        dup_resp = self.client.post(
            "/api/auth/register",
            json={"student_id": "20230901", "student_name": "Impostor", "password": "anotherPassword"},
        )
        self.assertEqual(dup_resp.status_code, 400)

    # =========================================================================
    # 8. FACULTY PORTAL HIDES PASSWORD & BLOCKS STUDENT / PI ACCESS
    # =========================================================================
    def test_08_faculty_portal_hides_password_and_blocks_student_and_pi(self):
        # Unauthenticated visit to /faculty shows login overlay WITHOUT exposing faculty2026 or pi2026 in HTML
        self.client.post("/api/faculty/logout")
        self.client.post("/api/auth/logout")
        fac_page = self.client.get("/faculty")
        html = fac_page.get_data(as_text=True)
        self.assertNotIn("e.g. faculty2026", html)
        self.assertNotIn("Passcode: pi2026", html)
        self.assertNotIn('id="btn-pi-supervisor-login"', html)
        self.assertNotIn('id="btn-show-passcode-from-denied"', html)

        # Student session attempting to visit /faculty gets 403 Forbidden
        self._login_student("20230001")
        student_fac_resp = self.client.get("/faculty")
        self.assertEqual(student_fac_resp.status_code, 403)
        self.assertIn('data-access-denied="true"', student_fac_resp.get_data(as_text=True))

        # PI session attempting to visit /faculty gets 403 Forbidden
        self._login_pi("PI-RANIA", "pi2026")
        pi_fac_resp = self.client.get("/faculty")
        self.assertEqual(pi_fac_resp.status_code, 403)
        self.assertIn('data-access-denied="true"', pi_fac_resp.get_data(as_text=True))


if __name__ == "__main__":
    unittest.main()
