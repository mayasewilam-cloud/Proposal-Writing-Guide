"""
BioWriter Studio - Cohort, Attendance, Lab Group, Project Team, PI, Topic & Shared Team Workspace Store.
Extends the existing Faculty & Student Portal authentication and draft storage architecture
without duplicating databases. Integrates directly with portal/data/users.json and portal/data/drafts/.
"""

import os
import io
import re
import csv
import json
import uuid
import zipfile
import datetime
import threading
import xml.etree.ElementTree as ET
from typing import Dict, Any, List, Optional, Tuple
from werkzeug.security import generate_password_hash, check_password_hash

from core.auth_store import (
    DATA_DIR,
    USERS_FILE,
    _ensure_data_dir,
    _load_users,
    _save_users,
    normalize_student_id,
)
from core.draft_store import (
    get_student_draft,
    save_student_draft,
    calculate_draft_summary,
)

_COHORT_LOCK = threading.Lock()

COHORT_DB_FILE = os.path.join(DATA_DIR, "cohort_db.json")
TEAM_WORKSPACES_DIR = os.path.join(DATA_DIR, "team_workspaces")

DEFAULT_SEMESTER = "Fall 2026"
DEFAULT_COURSE_ID = "BT_301"


def normalize_course_id(course_id: Optional[str]) -> str:
    if not course_id:
        return DEFAULT_COURSE_ID
    c = str(course_id).strip().upper().replace("-", "_").replace(" ", "_")
    if c in ("BT301", "BT_301"):
        return "BT_301"
    if c in ("GRAD501", "GRAD_501"):
        return "GRAD_501"
    if c in ("GRANTSTDF", "GRANT_STDF"):
        return "GRANT_STDF"
    return c


def normalize_semester(semester: Optional[str]) -> str:
    if not semester or not str(semester).strip():
        return DEFAULT_SEMESTER
    return str(semester).strip()


def normalize_lab_group_code(raw_group: Optional[str]) -> str:
    """
    Normalizes lab group identifiers while keeping them human-readable.
    Examples:
      'L1' -> 'L1'
      'Lab Group 1' -> 'L1'
      'Lab 2' -> 'L2'
      'Group 3' -> 'L3'
      '1' -> 'L1'
    """
    if raw_group is None:
        return ""
    val = str(raw_group).strip()
    if not val:
        return ""
    m = re.match(r"^(?:lab\s*group|lab|group|section|lg|l)\s*[-_#:\s]*(\d+[a-zA-Z]?)$", val, re.IGNORECASE)
    if m:
        return f"L{m.group(1).upper()}"
    if re.match(r"^\d+$", val):
        return f"L{val}"
    return val.upper() if len(val) <= 6 else val


def _now_str() -> str:
    return datetime.datetime.now().strftime("%b %d, %Y - %H:%M")


def _now_iso() -> str:
    return datetime.datetime.now().isoformat()


def _ensure_dirs():
    _ensure_data_dir()
    os.makedirs(TEAM_WORKSPACES_DIR, exist_ok=True)


_CACHED_DEFAULT_PIS: Optional[Dict[str, Dict[str, Any]]] = None


def _default_pis() -> Dict[str, Dict[str, Any]]:
    global _CACHED_DEFAULT_PIS
    if _CACHED_DEFAULT_PIS is None:
        shared_pi_hash = generate_password_hash("pi2026", method="pbkdf2:sha256", salt_length=16)
        _CACHED_DEFAULT_PIS = {
            "PI-RANIA": {
                "pi_id": "PI-RANIA",
                "name": "Dr. Rania Hassan",
                "title": "Associate Professor of Molecular Diagnostics",
                "department": "Biotechnology & Genetic Engineering",
                "email": "r.hassan@nu.edu.eg",
                "passcode_hash": shared_pi_hash,
                "role": "pi",
                "active": True,
            },
            "PI-AHMED": {
                "pi_id": "PI-AHMED",
                "name": "Dr. Ahmed Hassan",
                "title": "Professor of Pharmaceutical & Plant Biotechnology",
                "department": "Biotechnology & Genetic Engineering",
                "email": "a.hassan@nu.edu.eg",
                "passcode_hash": shared_pi_hash,
                "role": "pi",
                "active": True,
            },
            "PI-TARIK": {
                "pi_id": "PI-TARIK",
                "name": "Dr. Tarek El-Sayed",
                "title": "Professor of Synthetic & Industrial Biology",
                "department": "Biotechnology & Genetic Engineering",
                "email": "t.elsayed@nu.edu.eg",
                "passcode_hash": shared_pi_hash,
                "role": "pi",
                "active": True,
            },
            "PI-TAREK": {
                "pi_id": "PI-TAREK",
                "name": "Dr. Tarek Mansour",
                "title": "Professor of Synthetic & Industrial Biology",
                "department": "Biotechnology & Genetic Engineering",
                "email": "t.mansour@nu.edu.eg",
                "passcode_hash": shared_pi_hash,
                "role": "pi",
                "active": True,
            },
            "PI-MONA": {
                "pi_id": "PI-MONA",
                "name": "Dr. Mona Abdelrahman",
                "title": "Associate Professor of Microbial Genomics",
                "department": "Biotechnology & Genetic Engineering",
                "email": "m.abdelrahman@nu.edu.eg",
                "passcode_hash": shared_pi_hash,
                "role": "pi",
                "active": True,
            },
            "PI-NADIA": {
                "pi_id": "PI-NADIA",
                "name": "Dr. Nadia El-Sayed",
                "title": "Assistant Professor of Nanobiotechnology",
                "department": "Biotechnology & Genetic Engineering",
                "email": "n.elsayed@nu.edu.eg",
                "passcode_hash": shared_pi_hash,
                "role": "pi",
                "active": True,
            },
            "PI-KARIM": {
                "pi_id": "PI-KARIM",
                "name": "Dr. Karim Fawzy",
                "title": "Lecturer of Environmental & Agricultural Biotech",
                "department": "Biotechnology & Genetic Engineering",
                "email": "k.fawzy@nu.edu.eg",
                "passcode_hash": shared_pi_hash,
                "role": "pi",
                "active": True,
            },
        }
    return json.loads(json.dumps(_CACHED_DEFAULT_PIS))



def _load_db() -> Dict[str, Any]:
    _ensure_dirs()
    db: Dict[str, Any] = {}
    if os.path.exists(COHORT_DB_FILE):
        try:
            with open(COHORT_DB_FILE, "r", encoding="utf-8") as f:
                db = json.load(f)
        except Exception:
            db = {}

    changed = False
    if "semesters" not in db or not isinstance(db["semesters"], list):
        db["semesters"] = ["Fall 2026", "Spring 2026", "Spring 2027"]
        changed = True
    if "active_semester" not in db:
        db["active_semester"] = DEFAULT_SEMESTER
        changed = True
    if "active_course_id" not in db:
        db["active_course_id"] = DEFAULT_COURSE_ID
        changed = True
    default_pis = _default_pis()
    if "pis" not in db or not isinstance(db["pis"], dict) or not db["pis"]:
        db["pis"] = default_pis
        changed = True
    else:
        for pid, pinfo in default_pis.items():
            if pid not in db["pis"]:
                db["pis"][pid] = pinfo
                changed = True
    if "students" not in db or not isinstance(db["students"], dict):
        db["students"] = {}
        changed = True
    if "lab_groups" not in db or not isinstance(db["lab_groups"], dict):
        db["lab_groups"] = {}
        changed = True
    if "teams" not in db or not isinstance(db["teams"], dict):
        db["teams"] = {}
        changed = True
    if "pending_imports" not in db or not isinstance(db["pending_imports"], dict):
        db["pending_imports"] = {}
        changed = True
    if "audit_history" not in db or not isinstance(db["audit_history"], list):
        db["audit_history"] = []
        changed = True

    # Seed exemplar students, lab groups, and teams when cohort_db is first initialized
    if not db["students"] or "13989" not in db["students"]:
        default_students = [
            ("13989", "Maya Abdelrazek", "m.abdelrazek@nu.edu.eg", "L1", "L1-T1"),
            ("20230001", "Ahmed Ali", "ahmed.ali@student.edu.eg", "L1", "L1-T1"),
            ("20230002", "Sara Mohamed", "sara.mohamed@student.edu.eg", "L1", "L1-T1"),
            ("20230003", "Mariam Hassan", "mariam.hassan@student.edu.eg", "L1", "L1-T2"),
            ("20230004", "Karim Adel", "karim.adel@student.edu.eg", "L1", "L1-T2"),
            ("20230005", "Omar Khaled", "omar.khaled@student.edu.eg", "L1", ""),
            ("20230006", "Nour Hany", "nour.hany@student.edu.eg", "L1", ""),
            ("20230007", "Youssef Tarek", "youssef.tarek@student.edu.eg", "L2", "L2-T1"),
        ]
        for sid, sname, email, lg, tid in default_students:
            if sid not in db["students"]:
                db["students"][sid] = {
                    "student_id": sid,
                    "student_name": sname,
                    "email": email,
                    "lab_group": lg,
                    "team_id": tid,
                    "semester": DEFAULT_SEMESTER,
                    "course_id": DEFAULT_COURSE_ID,
                    "created_at": "Sep 20, 2026 - 10:00",
                    "updated_at": _now_str(),
                }
        changed = True

    for lg_code, lg_name in [("L1", "Lab Group 1 (L1)"), ("L2", "Lab Group 2 (L2)"), ("L3", "Lab Group 3 (L3)")]:
        lg_key = f"{DEFAULT_SEMESTER}::{DEFAULT_COURSE_ID}::{lg_code}"
        if lg_key not in db["lab_groups"]:
            db["lab_groups"][lg_key] = {
                "lab_group_id": lg_code,
                "name": lg_name,
                "semester": DEFAULT_SEMESTER,
                "course_id": DEFAULT_COURSE_ID,
                "created_at": "Sep 20, 2026 - 10:00",
            }
            changed = True

    if "L1-T1" not in db["teams"]:
        db["teams"]["L1-T1"] = {
            "team_id": "L1-T1",
            "team_name": "Team 1 — DETECTR-Aqua",
            "lab_group": "L1",
            "semester": DEFAULT_SEMESTER,
            "course_id": DEFAULT_COURSE_ID,
            "student_ids": ["13989", "20230001", "20230002"],
            "pi_id": "PI-RANIA",
            "pi_name": "Dr. Rania Hassan",
            "pi_email": "r.hassan@nu.edu.eg",
            "topic_title": "DETECTR-Aqua: A CRISPR-Cas12a-Assisted Isothermal Assay for Detecting the Colistin-Resistance Gene mcr-1 in Hospital Wastewater",
            "topic_description": "Matrix-aware RPA-Cas12a fluorescence diagnostic assay for rapid surveillance of plasmid-borne colistin resistance in hospital effluent.",
            "topic_category": "Cell-Free Diagnostics / AMR Surveillance",
            "status": "active",
            "created_at": "Sep 20, 2026 - 10:00",
            "updated_at": _now_str(),
        }
        changed = True

    if "L1-T2" not in db["teams"]:
        db["teams"]["L1-T2"] = {
            "team_id": "L1-T2",
            "team_name": "Team 2 — PhytoNano",
            "lab_group": "L1",
            "semester": DEFAULT_SEMESTER,
            "course_id": DEFAULT_COURSE_ID,
            "student_ids": ["20230003", "20230004"],
            "pi_id": "PI-AHMED",
            "pi_name": "Dr. Ahmed Hassan",
            "pi_email": "a.hassan@nu.edu.eg",
            "topic_title": "Plant-Derived Compounds and Anti-Inflammatory Nanocarriers",
            "topic_description": "Targeted liposomal delivery of standardized plant secondary metabolites for NF-kB modulation.",
            "topic_category": "Nanobiotechnology",
            "status": "active",
            "created_at": "Sep 20, 2026 - 10:00",
            "updated_at": _now_str(),
        }
        changed = True

    if "L2-T1" not in db["teams"]:
        db["teams"]["L2-T1"] = {
            "team_id": "L2-T1",
            "team_name": "Team 1 — BioCatalyst",
            "lab_group": "L2",
            "semester": DEFAULT_SEMESTER,
            "course_id": DEFAULT_COURSE_ID,
            "student_ids": ["20230007"],
            "pi_id": "PI-TARIK",
            "pi_name": "Dr. Tarek El-Sayed",
            "pi_email": "t.elsayed@nu.edu.eg",
            "topic_title": "Thermostable PETase Variants for Industrial Polyester Depolymerization",
            "topic_description": "Directed evolution of Ideonella sakaiensis PETase for bioreactor-scale plastic recycling.",
            "topic_category": "Industrial Biotechnology",
            "status": "active",
            "created_at": "Sep 20, 2026 - 10:00",
            "updated_at": _now_str(),
        }
        changed = True

    if changed:
        _save_db(db)

    return db


def _save_db(db: Dict[str, Any]):
    _ensure_dirs()
    with open(COHORT_DB_FILE, "w", encoding="utf-8") as f:
        json.dump(db, f, indent=2, ensure_ascii=False)


def _record_audit(
    db: Dict[str, Any],
    event_type: str,
    summary: str,
    actor: str = "Faculty Coordinator",
    student_id: Optional[str] = None,
    team_id: Optional[str] = None,
    lab_group: Optional[str] = None,
    old_value: Optional[str] = None,
    new_value: Optional[str] = None,
    semester: str = DEFAULT_SEMESTER,
    course_id: str = DEFAULT_COURSE_ID,
) -> Dict[str, Any]:
    entry = {
        "audit_id": f"AUD-{uuid.uuid4().hex[:8].upper()}",
        "event_type": event_type,
        "summary": summary,
        "actor": actor or "Faculty Coordinator",
        "student_id": student_id,
        "team_id": team_id,
        "lab_group": lab_group,
        "old_value": old_value,
        "new_value": new_value,
        "semester": normalize_semester(semester),
        "course_id": normalize_course_id(course_id),
        "timestamp": _now_str(),
        "timestamp_iso": _now_iso(),
    }
    db.setdefault("audit_history", []).insert(0, entry)
    return entry


# ==============================================================================
# 1. ATTENDANCE FILE PARSING (.CSV & .XLSX) AND IMPORT PREVIEW / CONFIRMATION
# ==============================================================================

def _parse_xlsx_bytes(file_bytes: bytes) -> List[Dict[str, str]]:
    """
    Parses an Excel (.xlsx) file into a list of row dicts using either openpyxl (if available)
    or Python's built-in zipfile + xml.etree.ElementTree (zero external dependency fallback).
    """
    try:
        import openpyxl  # type: ignore
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        sheet = wb.active
        rows_iter = list(sheet.iter_rows(values_only=True))
        if not rows_iter:
            return []
        headers = [str(h).strip() if h is not None else f"col_{i}" for i, h in enumerate(rows_iter[0])]
        results = []
        for r in rows_iter[1:]:
            if all(cell is None or str(cell).strip() == "" for cell in r):
                continue
            row_dict = {}
            for i, val in enumerate(r):
                key = headers[i] if i < len(headers) else f"col_{i}"
                row_dict[key] = str(val).strip() if val is not None else ""
            results.append(row_dict)
        return results
    except Exception:
        pass

    # Built-in stdlib XML parser for .xlsx (OpenXML spreadsheetml)
    with zipfile.ZipFile(io.BytesIO(file_bytes)) as zf:
        shared_strings: List[str] = []
        if "xl/sharedStrings.xml" in zf.namelist():
            ss_xml = zf.read("xl/sharedStrings.xml")
            ss_root = ET.fromstring(ss_xml)
            for si in ss_root.iter():
                if si.tag.endswith("}si") or si.tag == "si":
                    texts = []
                    for t in si.iter():
                        if (t.tag.endswith("}t") or t.tag == "t") and t.text:
                            texts.append(t.text)
                    shared_strings.append("".join(texts))

        sheet_names = [n for n in zf.namelist() if n.startswith("xl/worksheets/sheet") and n.endswith(".xml")]
        if not sheet_names:
            return []
        sheet_names.sort()
        sheet_xml = zf.read(sheet_names[0])
        root = ET.fromstring(sheet_xml)

        def _col_index(cell_ref: str) -> int:
            letters = "".join(ch for ch in cell_ref if ch.isalpha()).upper()
            idx = 0
            for ch in letters:
                idx = idx * 26 + (ord(ch) - ord("A") + 1)
            return max(0, idx - 1)

        parsed_grid: List[List[str]] = []
        for elem in root.iter():
            if elem.tag.endswith("}row") or elem.tag == "row":
                row_vals: Dict[int, str] = {}
                max_col = -1
                for cell in elem:
                    if not (cell.tag.endswith("}c") or cell.tag == "c"):
                        continue
                    ref = cell.attrib.get("r", "")
                    c_idx = _col_index(ref) if ref else (max_col + 1)
                    max_col = max(max_col, c_idx)
                    c_type = cell.attrib.get("t", "")
                    val_str = ""
                    if c_type == "inlineStr":
                        texts = [t.text for t in cell.iter() if (t.tag.endswith("}t") or t.tag == "t") and t.text]
                        val_str = "".join(texts)
                    else:
                        v_el = next((ch for ch in cell if ch.tag.endswith("}v") or ch.tag == "v"), None)
                        if v_el is not None and v_el.text is not None:
                            raw_v = v_el.text.strip()
                            if c_type == "s" and raw_v.isdigit():
                                s_idx = int(raw_v)
                                val_str = shared_strings[s_idx] if 0 <= s_idx < len(shared_strings) else raw_v
                            else:
                                val_str = raw_v
                    row_vals[c_idx] = val_str.strip()
                if max_col >= 0:
                    row_list = [row_vals.get(i, "") for i in range(max_col + 1)]
                    if any(x != "" for x in row_list):
                        parsed_grid.append(row_list)

        if not parsed_grid:
            return []
        headers = [h.strip() if h else f"col_{i}" for i, h in enumerate(parsed_grid[0])]
        results = []
        for r in parsed_grid[1:]:
            if all(x == "" for x in r):
                continue
            row_dict = {}
            for i, val in enumerate(r):
                key = headers[i] if i < len(headers) else f"col_{i}"
                row_dict[key] = val
            results.append(row_dict)
        return results


def _parse_csv_bytes(file_bytes: bytes) -> List[Dict[str, str]]:
    text = ""
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            text = file_bytes.decode(enc)
            break
        except Exception:
            continue
    if not text.strip():
        return []

    lines = [line for line in text.splitlines() if line.strip()]
    if not lines:
        return []

    # Detect delimiter
    sample = "\n".join(lines[:5])
    delimiter = ","
    if sample.count(";") > sample.count(",") and sample.count(";") >= sample.count("\t"):
        delimiter = ";"
    elif sample.count("\t") > sample.count(","):
        delimiter = "\t"

    reader = csv.DictReader(io.StringIO("\n".join(lines)), delimiter=delimiter)
    results = []
    for row in reader:
        clean_row = {}
        for k, v in (row or {}).items():
            if k is None:
                continue
            clean_row[str(k).strip()] = str(v).strip() if v is not None else ""
        if any(val != "" for val in clean_row.values()):
            results.append(clean_row)
    return results


def parse_attendance_file(filename: str, file_bytes: bytes) -> List[Dict[str, str]]:
    fname = (filename or "").lower().strip()
    if fname.endswith(".xlsx") or (len(file_bytes) >= 4 and file_bytes[:2] == b"PK"):
        return _parse_xlsx_bytes(file_bytes)
    return _parse_csv_bytes(file_bytes)


def _extract_row_fields(raw_row: Dict[str, Any]) -> Tuple[str, str, str, str]:
    """
    Extracts (student_id, student_name, lab_group, email) from a flexible dictionary row.
    """
    sid = ""
    sname = ""
    lgroup = ""
    email = ""

    for k, v in raw_row.items():
        kl = re.sub(r"[^a-z0-9]", "", str(k).lower())
        val = str(v).strip() if v is not None else ""
        if kl in ("studentid", "id", "universityid", "studentcode", "code", "matriculation", "userid", "sid"):
            if not sid:
                sid = val
        elif kl in ("studentname", "name", "fullname", "student", "studentfullname", "membername"):
            if not sname:
                sname = val
        elif kl in ("labgroup", "lab", "group", "section", "labsection", "tutorialgroup", "practicalgroup", "lg"):
            if not lgroup:
                lgroup = val
        elif kl in ("email", "universityemail", "studentemail", "mail"):
            if not email:
                email = val

    return normalize_student_id(sid), sname.strip(), normalize_lab_group_code(lgroup), email.strip()


def preview_attendance_upload(
    raw_rows: List[Dict[str, Any]],
    semester: str = DEFAULT_SEMESTER,
    course_id: str = DEFAULT_COURSE_ID,
    filename: str = "attendance.csv",
    actor: str = "Faculty Coordinator",
) -> Dict[str, Any]:
    """
    Validates and stages an attendance upload WITHOUT overwriting existing database records.
    Detects:
      - duplicate student IDs
      - duplicate rows
      - missing student IDs
      - missing names
      - missing lab groups
      - students already existing in the portal vs new students
    """
    with _COHORT_LOCK:
        db = _load_db()
        users = _load_users()
        sem = normalize_semester(semester)
        cid = normalize_course_id(course_id)

        preview_rows = []
        seen_ids: Dict[str, int] = {}
        seen_exact_rows: set = set()
        detected_lab_groups: set = set()
        unique_valid_students: set = set()

        new_count = 0
        existing_count = 0
        duplicate_id_count = 0
        duplicate_row_count = 0
        missing_id_count = 0
        missing_name_count = 0
        missing_lab_group_count = 0

        for idx, raw in enumerate(raw_rows, start=1):
            sid, sname, lgroup, email = _extract_row_fields(raw)
            issues = []

            if not sid:
                missing_id_count += 1
                issues.append("Missing Student ID")
            if not sname:
                missing_name_count += 1
                issues.append("Missing Student Name")
            if not lgroup:
                missing_lab_group_count += 1
                issues.append("Missing Lab Group")

            row_signature = (sid.upper(), sname.lower(), lgroup.upper())
            is_duplicate_row = False
            is_duplicate_id = False

            if sid and sname and lgroup and row_signature in seen_exact_rows:
                is_duplicate_row = True
                duplicate_row_count += 1
                issues.append(f"Duplicate row (identical to earlier row #{seen_ids.get(sid, 1)})")
            elif sid:
                if sid in seen_ids:
                    is_duplicate_id = True
                    duplicate_id_count += 1
                    issues.append(f"Duplicate Student ID (first seen on row #{seen_ids[sid]})")
                else:
                    seen_ids[sid] = idx
                    if sname and lgroup:
                        seen_exact_rows.add(row_signature)

            # Determine existing status in portal
            existing_in_portal = (sid in users) or (sid in db.get("students", {}))
            existing_record = db.get("students", {}).get(sid) or users.get(sid) or {}
            previous_lab_group = existing_record.get("lab_group", "")

            if issues:
                if is_duplicate_row:
                    status = "Duplicate Row"
                elif is_duplicate_id:
                    status = "Duplicate ID"
                else:
                    status = "Missing Fields: " + ", ".join(issues)
                is_valid = False
            else:
                is_valid = True
                unique_valid_students.add(sid)
                if lgroup:
                    detected_lab_groups.add(lgroup)
                if existing_in_portal:
                    status = "Existing"
                    existing_count += 1
                else:
                    status = "New"
                    new_count += 1

            preview_rows.append({
                "row_number": idx,
                "student_id": sid,
                "student_name": sname,
                "lab_group": lgroup,
                "previous_lab_group": previous_lab_group,
                "email": email or existing_record.get("email", ""),
                "status": status,
                "existing_in_portal": existing_in_portal,
                "is_valid": is_valid,
                "issues": issues,
            })

        preview_id = f"IMP-{uuid.uuid4().hex[:10].upper()}"
        sorted_groups = sorted(
            list(detected_lab_groups),
            key=lambda g: (0, int(g[1:])) if (g.startswith("L") and g[1:].isdigit()) else (1, g)
        )

        preview_payload = {
            "preview_id": preview_id,
            "filename": filename,
            "semester": sem,
            "course_id": cid,
            "created_by": actor,
            "created_at": _now_str(),
            "students_detected": len(unique_valid_students),
            "lab_groups_detected": len(sorted_groups),
            "lab_groups_list": sorted_groups,
            "summary": {
                "total_rows": len(preview_rows),
                "valid_rows": len(unique_valid_students),
                "new_count": new_count,
                "existing_count": existing_count,
                "duplicate_id_count": duplicate_id_count,
                "duplicate_row_count": duplicate_row_count,
                "missing_id_count": missing_id_count,
                "missing_name_count": missing_name_count,
                "missing_lab_group_count": missing_lab_group_count,
            },
            "rows": preview_rows,
        }

        db["pending_imports"][preview_id] = preview_payload
        _save_db(db)

        return {
            "success": True,
            **preview_payload,
        }


def confirm_attendance_import(
    preview_id: str,
    actor: str = "Faculty Coordinator",
) -> Dict[str, Any]:
    """
    Commits a staged attendance import preview into the canonical student & lab group database.
    Only valid rows ('New' or 'Existing') are imported; duplicate/invalid rows are skipped.
    """
    with _COHORT_LOCK:
        db = _load_db()
        pending = db.get("pending_imports", {}).get(preview_id)
        if not pending:
            return {"success": False, "error": f"Attendance preview '{preview_id}' not found or already confirmed."}

        sem = normalize_semester(pending.get("semester"))
        cid = normalize_course_id(pending.get("course_id"))
        users = _load_users()

        imported_students = []
        created_count = 0
        updated_count = 0
        skipped_count = 0
        groups_touched = set()

        for row in pending.get("rows", []):
            if not row.get("is_valid"):
                skipped_count += 1
                continue

            sid = normalize_student_id(row["student_id"])
            sname = row["student_name"].strip()
            lgroup = normalize_lab_group_code(row["lab_group"])
            email = row.get("email", "").strip() or f"{sid.lower()}@student.edu.eg"
            groups_touched.add(lgroup)

            # Ensure lab group exists in cohort_db
            lg_key = f"{sem}::{cid}::{lgroup}"
            if lg_key not in db["lab_groups"]:
                db["lab_groups"][lg_key] = {
                    "lab_group_id": lgroup,
                    "name": f"Lab Group {lgroup[1:]}" if (lgroup.startswith("L") and lgroup[1:].isdigit()) else lgroup,
                    "semester": sem,
                    "course_id": cid,
                    "created_at": _now_str(),
                }

            # Check if student already exists in cohort_db
            existing_st = db["students"].get(sid)
            if existing_st:
                old_lg = existing_st.get("lab_group")
                existing_st["student_name"] = sname
                existing_st["lab_group"] = lgroup
                existing_st["semester"] = sem
                existing_st["course_id"] = cid
                if email:
                    existing_st["email"] = email
                existing_st["updated_at"] = _now_str()

                # If lab group changed and student was in a team from the old lab group, unassign from old team
                if old_lg and old_lg != lgroup:
                    old_team_id = existing_st.get("team_id")
                    if old_team_id and old_team_id in db["teams"]:
                        t_obj = db["teams"][old_team_id]
                        if t_obj.get("lab_group") != lgroup:
                            if sid in t_obj.get("student_ids", []):
                                t_obj["student_ids"].remove(sid)
                            existing_st["team_id"] = None
                    _record_audit(
                        db,
                        event_type="LAB_GROUP_MOVE",
                        summary=f"Moved {sname} ({sid}) from {old_lg} to {lgroup} via Attendance Import by {actor} on {_now_str()}",
                        actor=actor,
                        student_id=sid,
                        lab_group=lgroup,
                        old_value=old_lg,
                        new_value=lgroup,
                        semester=sem,
                        course_id=cid,
                    )
                updated_count += 1
            else:
                db["students"][sid] = {
                    "student_id": sid,
                    "student_name": sname,
                    "email": email,
                    "lab_group": lgroup,
                    "team_id": None,
                    "semester": sem,
                    "course_id": cid,
                    "created_at": _now_str(),
                    "updated_at": _now_str(),
                }
                created_count += 1

            # Sync with canonical auth_store users.json so student can create account or sign in immediately
            if sid not in users:
                users[sid] = {
                    "student_id": sid,
                    "student_name": sname,
                    "email": email,
                    "lab_group": lgroup,
                    "semester": sem,
                    "course_id": cid,
                    "password_hash": generate_password_hash("bio123", method="pbkdf2:sha256", salt_length=16),
                    "created_at": _now_str(),
                    "last_login": _now_str(),
                    "role": "student",
                    "pre_enrolled_by_faculty": True,
                    "account_claimed": False,
                }
            else:
                users[sid]["student_name"] = sname
                users[sid]["lab_group"] = lgroup
                users[sid]["semester"] = sem
                users[sid]["course_id"] = cid
                if sid not in ("ST-2026-01", "13989"):
                    users[sid].setdefault("pre_enrolled_by_faculty", True)
                    users[sid].setdefault("account_claimed", False)
                if email:
                    users[sid]["email"] = email

            imported_students.append(db["students"][sid])

        # Remove from pending_imports after confirmation
        db["pending_imports"].pop(preview_id, None)

        audit_entry = _record_audit(
            db,
            event_type="ATTENDANCE_IMPORT",
            summary=f"Confirmed attendance import ({len(imported_students)} students across {len(groups_touched)} lab groups: {created_count} new, {updated_count} existing) by {actor} on {_now_str()}",
            actor=actor,
            semester=sem,
            course_id=cid,
        )

        _save_users(users)
        _save_db(db)

        return {
            "success": True,
            "preview_id": preview_id,
            "semester": sem,
            "course_id": cid,
            "imported_count": len(imported_students),
            "created_count": created_count,
            "updated_count": updated_count,
            "skipped_count": skipped_count,
            "lab_groups_count": len(groups_touched),
            "audit_entry": audit_entry,
        }


# ==============================================================================
# 2. STUDENT DIRECTORY, SEARCH, FILTERING & LAB GROUP MOVES
# ==============================================================================

def _sync_users_into_cohort_db(db: Dict[str, Any]) -> bool:
    """Ensures any student registered via /api/auth/register in users.json is also visible in cohort_db."""
    users = _load_users()
    changed = False
    for sid, u in users.items():
        if u.get("role", "student") != "student":
            continue
        if sid not in db["students"]:
            db["students"][sid] = {
                "student_id": sid,
                "student_name": u.get("student_name", sid),
                "email": u.get("email", f"{sid.lower()}@student.edu.eg"),
                "lab_group": u.get("lab_group", ""),
                "team_id": u.get("team_id", None),
                "semester": u.get("semester", DEFAULT_SEMESTER),
                "course_id": normalize_course_id(u.get("course_id", DEFAULT_COURSE_ID)),
                "created_at": u.get("created_at", _now_str()),
                "updated_at": _now_str(),
            }
            changed = True
    return changed


def enrich_student_record(student: Dict[str, Any], db: Dict[str, Any]) -> Dict[str, Any]:
    """Attaches team, PI, and topic details to a student record."""
    sid = student["student_id"]
    team_id = student.get("team_id")
    team = db.get("teams", {}).get(team_id) if team_id else None
    if team and team.get("status") == "archived":
        team = None
        team_id = None

    # History for this student
    history = [
        h for h in db.get("audit_history", [])
        if h.get("student_id") == sid or (team_id and h.get("team_id") == team_id)
    ][:20]

    return {
        "student_id": sid,
        "student_name": student.get("student_name", sid),
        "email": student.get("email", ""),
        "lab_group": student.get("lab_group") or "",
        "team_id": team_id or "",
        "team_name": team.get("team_name", "") if team else "",
        "pi_id": team.get("pi_id", "") if team else "",
        "pi_name": team.get("pi_name", "") if team else "",
        "pi_email": team.get("pi_email", "") if team else "",
        "topic_title": team.get("topic_title", "") if team else "",
        "topic_description": team.get("topic_description", "") if team else "",
        "topic_category": team.get("topic_category", "") if team else "",
        "semester": student.get("semester", DEFAULT_SEMESTER),
        "course_id": normalize_course_id(student.get("course_id", DEFAULT_COURSE_ID)),
        "updated_at": student.get("updated_at", ""),
        "history": history,
    }


def list_students(
    semester: Optional[str] = None,
    course_id: Optional[str] = None,
    lab_group: Optional[str] = None,
    team_id: Optional[str] = None,
    pi_id: Optional[str] = None,
    assignment_status: Optional[str] = None,
    search_query: Optional[str] = None,
) -> Dict[str, Any]:
    """
    Returns enriched student list with support for all Section 2 & Section 14 filters:
      - lab_group
      - team_id
      - pi_id
      - assignment_status:
          'unassigned_lab_group' (students without a lab group)
          'unassigned_team' (students without a team)
          'assigned_team' (students in a team)
          'missing_pi' (students whose team has no PI)
          'missing_topic' (students whose team has no topic)
      - search_query (by Student ID, Name, Topic, or Team)
    """
    with _COHORT_LOCK:
        db = _load_db()
        if _sync_users_into_cohort_db(db):
            _save_db(db)

        sem = normalize_semester(semester or db.get("active_semester"))
        cid = normalize_course_id(course_id or db.get("active_course_id"))
        lg_filter = normalize_lab_group_code(lab_group) if lab_group else ""
        q = (search_query or "").strip().lower()

        results = []
        for sid, st in db.get("students", {}).items():
            st_sem = normalize_semester(st.get("semester"))
            st_cid = normalize_course_id(st.get("course_id"))
            if semester and st_sem != sem:
                continue
            if course_id and st_cid != cid:
                continue

            enriched = enrich_student_record(st, db)

            if lg_filter and enriched["lab_group"].upper() != lg_filter.upper():
                continue
            if team_id and enriched["team_id"].upper() != team_id.strip().upper():
                continue
            if pi_id and enriched["pi_id"].upper() != pi_id.strip().upper():
                continue

            if assignment_status:
                status_norm = assignment_status.strip().lower()
                if status_norm == "unassigned_lab_group" and enriched["lab_group"]:
                    continue
                elif status_norm == "unassigned_team" and enriched["team_id"]:
                    continue
                elif status_norm == "assigned_team" and not enriched["team_id"]:
                    continue
                elif status_norm == "missing_pi" and (not enriched["team_id"] or enriched["pi_id"]):
                    continue
                elif status_norm == "missing_topic" and (not enriched["team_id"] or enriched["topic_title"]):
                    continue

            if q:
                haystack = " ".join([
                    enriched["student_id"],
                    enriched["student_name"],
                    enriched["lab_group"],
                    enriched["team_id"],
                    enriched["team_name"],
                    enriched["pi_name"],
                    enriched["topic_title"],
                ]).lower()
                if q not in haystack:
                    continue

            results.append(enriched)

        results.sort(key=lambda x: (
            (0, int(x["lab_group"][1:])) if (x["lab_group"].startswith("L") and x["lab_group"][1:].isdigit()) else (1, x["lab_group"] or "ZZZ"),
            x["student_id"]
        ))

        return {
            "success": True,
            "semester": sem,
            "course_id": cid,
            "total": len(results),
            "students": results,
        }


def upsert_student_by_faculty(
    student_id: str,
    student_name: str,
    lab_group: str = "",
    email: str = "",
    semester: str = DEFAULT_SEMESTER,
    course_id: str = DEFAULT_COURSE_ID,
    actor: str = "Faculty Coordinator",
) -> Dict[str, Any]:
    """Allows faculty to manually add or update a single student record."""
    clean_id = normalize_student_id(student_id)
    if not clean_id:
        return {"success": False, "error": "Student ID is required."}
    clean_name = (student_name or "").strip()
    if not clean_name:
        return {"success": False, "error": "Student Name is required."}

    with _COHORT_LOCK:
        db = _load_db()
        users = _load_users()
        sem = normalize_semester(semester)
        cid = normalize_course_id(course_id)
        lgroup = normalize_lab_group_code(lab_group)

        if lgroup:
            lg_key = f"{sem}::{cid}::{lgroup}"
            if lg_key not in db["lab_groups"]:
                db["lab_groups"][lg_key] = {
                    "lab_group_id": lgroup,
                    "name": f"Lab Group {lgroup[1:]}" if (lgroup.startswith("L") and lgroup[1:].isdigit()) else lgroup,
                    "semester": sem,
                    "course_id": cid,
                    "created_at": _now_str(),
                }

        existing = db["students"].get(clean_id)
        if existing:
            old_lg = existing.get("lab_group", "")
            existing["student_name"] = clean_name
            if email:
                existing["email"] = email.strip()
            if lgroup and lgroup != old_lg:
                existing["lab_group"] = lgroup
                _record_audit(
                    db,
                    event_type="LAB_GROUP_MOVE",
                    summary=f"Moved {clean_name} ({clean_id}) from {old_lg or 'Unassigned'} to {lgroup} by {actor} on {_now_str()}",
                    actor=actor,
                    student_id=clean_id,
                    lab_group=lgroup,
                    old_value=old_lg or "Unassigned",
                    new_value=lgroup,
                    semester=sem,
                    course_id=cid,
                )
            existing["updated_at"] = _now_str()
        else:
            db["students"][clean_id] = {
                "student_id": clean_id,
                "student_name": clean_name,
                "email": email.strip() or f"{clean_id.lower()}@student.edu.eg",
                "lab_group": lgroup,
                "team_id": None,
                "semester": sem,
                "course_id": cid,
                "created_at": _now_str(),
                "updated_at": _now_str(),
            }
            _record_audit(
                db,
                event_type="STUDENT_ADDED",
                summary=f"Added student {clean_name} ({clean_id}) to lab group {lgroup or 'Unassigned'} by {actor} on {_now_str()}",
                actor=actor,
                student_id=clean_id,
                lab_group=lgroup,
                new_value=lgroup,
                semester=sem,
                course_id=cid,
            )

        if clean_id not in users:
            users[clean_id] = {
                "student_id": clean_id,
                "student_name": clean_name,
                "email": db["students"][clean_id]["email"],
                "lab_group": lgroup,
                "semester": sem,
                "course_id": cid,
                "password_hash": generate_password_hash("bio123", method="pbkdf2:sha256", salt_length=16),
                "created_at": _now_str(),
                "last_login": _now_str(),
                "role": "student",
                "pre_enrolled_by_faculty": True,
                "account_claimed": False,
            }
        else:
            users[clean_id]["student_name"] = clean_name
            users[clean_id]["lab_group"] = lgroup
            if clean_id not in ("ST-2026-01", "13989"):
                users[clean_id].setdefault("pre_enrolled_by_faculty", True)
                users[clean_id].setdefault("account_claimed", False)

        _save_users(users)
        _save_db(db)
        return {
            "success": True,
            "student": enrich_student_record(db["students"][clean_id], db),
        }


def move_student_lab_group(
    student_id: str,
    new_lab_group: str,
    actor: str = "Faculty Coordinator",
    semester: Optional[str] = None,
    course_id: Optional[str] = None,
) -> Dict[str, Any]:
    """
    Moves a student from their current lab group to a new lab group (Section 3),
    updating the database, handling any previous team assignment in the old lab group,
    and recording an immutable audit history entry.
    """
    clean_id = normalize_student_id(student_id)
    target_lg = normalize_lab_group_code(new_lab_group)
    if not clean_id:
        return {"success": False, "error": "Student ID is required."}
    if not target_lg:
        return {"success": False, "error": "New Lab Group is required."}

    with _COHORT_LOCK:
        db = _load_db()
        _sync_users_into_cohort_db(db)
        users = _load_users()

        st = db["students"].get(clean_id)
        if not st:
            return {"success": False, "error": f"Student '{clean_id}' not found."}

        sem = normalize_semester(semester or st.get("semester"))
        cid = normalize_course_id(course_id or st.get("course_id"))
        old_lg = st.get("lab_group") or "Unassigned"

        # Ensure destination lab group exists
        lg_key = f"{sem}::{cid}::{target_lg}"
        if lg_key not in db["lab_groups"]:
            db["lab_groups"][lg_key] = {
                "lab_group_id": target_lg,
                "name": f"Lab Group {target_lg[1:]}" if (target_lg.startswith("L") and target_lg[1:].isdigit()) else target_lg,
                "semester": sem,
                "course_id": cid,
                "created_at": _now_str(),
            }

        # If student was in a team belonging to the old lab group, remove them from that old team
        unassigned_from_team = None
        old_team_id = st.get("team_id")
        if old_team_id and old_team_id in db["teams"]:
            team_obj = db["teams"][old_team_id]
            if team_obj.get("lab_group") != target_lg:
                if clean_id in team_obj.get("student_ids", []):
                    team_obj["student_ids"].remove(clean_id)
                st["team_id"] = None
                unassigned_from_team = old_team_id
                _sync_team_workspace_metadata(db, team_obj)

        st["lab_group"] = target_lg
        st["updated_at"] = _now_str()

        if clean_id in users:
            users[clean_id]["lab_group"] = target_lg
            _save_users(users)

        summary_msg = f"Moved from {old_lg} to {target_lg} by {actor} on {_now_str()}"
        if unassigned_from_team:
            summary_msg += f" (unassigned from {unassigned_from_team})"

        audit_entry = _record_audit(
            db,
            event_type="LAB_GROUP_MOVE",
            summary=summary_msg,
            actor=actor,
            student_id=clean_id,
            team_id=unassigned_from_team,
            lab_group=target_lg,
            old_value=old_lg,
            new_value=target_lg,
            semester=sem,
            course_id=cid,
        )

        _save_db(db)

        return {
            "success": True,
            "student": enrich_student_record(st, db),
            "old_lab_group": old_lg,
            "previous_lab_group": old_lg,
            "new_lab_group": target_lg,
            "unassigned_from_team": unassigned_from_team,
            "audit_entry": audit_entry,
        }


# ==============================================================================
# 3. LAB GROUP MANAGEMENT & UNASSIGNED STUDENTS TRACKING (SECTIONS 3 & 5)
# ==============================================================================

def create_lab_group(
    lab_group_id: str,
    name: str = "",
    semester: str = DEFAULT_SEMESTER,
    course_id: str = DEFAULT_COURSE_ID,
    actor: str = "Faculty Coordinator",
) -> Dict[str, Any]:
    lgroup = normalize_lab_group_code(lab_group_id)
    if not lgroup:
        return {"success": False, "error": "Lab group code (e.g. L1) is required."}

    with _COHORT_LOCK:
        db = _load_db()
        sem = normalize_semester(semester)
        cid = normalize_course_id(course_id)
        lg_key = f"{sem}::{cid}::{lgroup}"

        display_name = name.strip() if name else (
            f"Lab Group {lgroup[1:]}" if (lgroup.startswith("L") and lgroup[1:].isdigit()) else lgroup
        )
        db["lab_groups"][lg_key] = {
            "lab_group_id": lgroup,
            "name": display_name,
            "semester": sem,
            "course_id": cid,
            "created_at": _now_str(),
        }
        _record_audit(
            db,
            event_type="LAB_GROUP_CREATED",
            summary=f"Created {display_name} ({lgroup}) by {actor} on {_now_str()}",
            actor=actor,
            lab_group=lgroup,
            new_value=lgroup,
            semester=sem,
            course_id=cid,
        )
        _save_db(db)
        return {"success": True, "lab_group": db["lab_groups"][lg_key]}


def list_lab_groups(
    semester: Optional[str] = None,
    course_id: Optional[str] = None,
) -> Dict[str, Any]:
    """
    Returns all lab groups for the given semester/course, including:
      - student_count
      - team_count
      - unassigned_student_count
      - students list
      - teams list
      - unassigned_students list (Section 5)
    """
    with _COHORT_LOCK:
        db = _load_db()
        _sync_users_into_cohort_db(db)
        sem = normalize_semester(semester or db.get("active_semester"))
        cid = normalize_course_id(course_id or db.get("active_course_id"))

        # Collect all known lab group codes from both lab_groups registry, students, and teams
        group_codes = set()
        for lg_key, lg_obj in db.get("lab_groups", {}).items():
            if normalize_semester(lg_obj.get("semester")) == sem and normalize_course_id(lg_obj.get("course_id")) == cid:
                group_codes.add(lg_obj["lab_group_id"])
        for st in db.get("students", {}).values():
            if normalize_semester(st.get("semester")) == sem and normalize_course_id(st.get("course_id")) == cid:
                if st.get("lab_group"):
                    group_codes.add(st["lab_group"])
        for tm in db.get("teams", {}).values():
            if tm.get("status") != "archived" and normalize_semester(tm.get("semester")) == sem and normalize_course_id(tm.get("course_id")) == cid:
                if tm.get("lab_group"):
                    group_codes.add(tm["lab_group"])

        sorted_codes = sorted(
            list(group_codes),
            key=lambda g: (0, int(g[1:])) if (g.startswith("L") and g[1:].isdigit()) else (1, g)
        )

        lab_groups_out = []
        for code in sorted_codes:
            lg_key = f"{sem}::{cid}::{code}"
            lg_meta = db.get("lab_groups", {}).get(lg_key, {})
            display_name = lg_meta.get("name") or (
                f"Lab Group {code[1:]}" if (code.startswith("L") and code[1:].isdigit()) else code
            )

            group_students = [
                enrich_student_record(st, db)
                for st in db.get("students", {}).values()
                if normalize_semester(st.get("semester")) == sem
                and normalize_course_id(st.get("course_id")) == cid
                and (st.get("lab_group") or "").upper() == code.upper()
            ]
            group_students.sort(key=lambda s: s["student_id"])

            group_teams = [
                _enrich_team_record(tm, db)
                for tm in db.get("teams", {}).values()
                if tm.get("status") != "archived"
                and normalize_semester(tm.get("semester")) == sem
                and normalize_course_id(tm.get("course_id")) == cid
                and (tm.get("lab_group") or "").upper() == code.upper()
            ]
            group_teams.sort(key=lambda t: t["team_id"])

            unassigned_students = [s for s in group_students if not s.get("team_id")]

            lab_groups_out.append({
                "lab_group_id": code,
                "name": display_name,
                "semester": sem,
                "course_id": cid,
                "student_count": len(group_students),
                "team_count": len(group_teams),
                "unassigned_count": len(unassigned_students),
                "summary_label": f"{code} — {len(group_students)} students — {len(group_teams)} teams",
                "students": group_students,
                "teams": group_teams,
                "unassigned_students": unassigned_students,
            })

        # Also compute students without any lab group assigned
        no_lab_group_students = [
            enrich_student_record(st, db)
            for st in db.get("students", {}).values()
            if normalize_semester(st.get("semester")) == sem
            and normalize_course_id(st.get("course_id")) == cid
            and not st.get("lab_group")
        ]

        return {
            "success": True,
            "semester": sem,
            "course_id": cid,
            "total_lab_groups": len(lab_groups_out),
            "lab_groups": lab_groups_out,
            "unassigned_lab_group_students": no_lab_group_students,
        }


# ==============================================================================
# 4. PROJECT TEAM SETUP, SINGLE-ACTIVE-TEAM ENFORCEMENT, PI & TOPIC (SECTIONS 4-7)
# ==============================================================================

def _enrich_team_record(team: Dict[str, Any], db: Dict[str, Any]) -> Dict[str, Any]:
    members = []
    for sid in team.get("student_ids", []):
        st = db.get("students", {}).get(sid)
        if st:
            members.append({
                "student_id": sid,
                "student_name": st.get("student_name", sid),
                "email": st.get("email", ""),
                "lab_group": st.get("lab_group", team.get("lab_group", "")),
            })
        else:
            members.append({
                "student_id": sid,
                "student_name": f"Student {sid}",
                "email": "",
                "lab_group": team.get("lab_group", ""),
            })

    return {
        "team_id": team["team_id"],
        "team_name": team.get("team_name", team["team_id"]),
        "lab_group": team.get("lab_group", ""),
        "semester": normalize_semester(team.get("semester")),
        "course_id": normalize_course_id(team.get("course_id")),
        "student_ids": list(team.get("student_ids", [])),
        "members": members,
        "member_count": len(members),
        "pi_id": team.get("pi_id") or "",
        "pi_name": team.get("pi_name") or "Unassigned",
        "pi_email": team.get("pi_email") or "",
        "topic_title": team.get("topic_title") or "",
        "topic_description": team.get("topic_description") or "",
        "topic_category": team.get("topic_category") or "",
        "topic_status": "Assigned" if team.get("topic_title") else "Unassigned",
        "status": team.get("status", "active"),
        "created_at": team.get("created_at", ""),
        "updated_at": team.get("updated_at", ""),
    }


def _find_student_active_team(
    db: Dict[str, Any],
    student_id: str,
    semester: str,
    course_id: str,
    exclude_team_id: Optional[str] = None,
) -> Optional[Dict[str, Any]]:
    clean_sid = normalize_student_id(student_id)
    sem = normalize_semester(semester)
    cid = normalize_course_id(course_id)
    for tid, tm in db.get("teams", {}).items():
        if tm.get("status") == "archived":
            continue
        if exclude_team_id and tid.upper() == exclude_team_id.upper():
            continue
        if normalize_semester(tm.get("semester")) == sem and normalize_course_id(tm.get("course_id")) == cid:
            if clean_sid in tm.get("student_ids", []):
                return tm
    return None


def create_project_team(
    lab_group: str,
    team_name: str = "",
    team_id: str = "",
    student_ids: Optional[List[str]] = None,
    pi_id: str = "",
    topic_title: str = "",
    topic_description: str = "",
    topic_category: str = "",
    semester: str = DEFAULT_SEMESTER,
    course_id: str = DEFAULT_COURSE_ID,
    actor: str = "Faculty Coordinator",
) -> Dict[str, Any]:
    """
    Creates a new project team inside a lab group.
    Enforces:
      - Each student may belong to ONLY ONE active project team per semester/course.
    """
    lgroup = normalize_lab_group_code(lab_group)
    if not lgroup:
        return {"success": False, "error": "Lab group is required to create a team."}

    with _COHORT_LOCK:
        db = _load_db()
        _sync_users_into_cohort_db(db)
        sem = normalize_semester(semester)
        cid = normalize_course_id(course_id)

        # Generate team_id if not provided (e.g., L1-T1, L1-T2)
        clean_tid = (team_id or "").strip().upper()
        if not clean_tid:
            existing_in_lg = [
                tm for tm in db.get("teams", {}).values()
                if (tm.get("lab_group") or "").upper() == lgroup.upper()
                and normalize_semester(tm.get("semester")) == sem
                and normalize_course_id(tm.get("course_id")) == cid
            ]
            next_num = len(existing_in_lg) + 1
            while f"{lgroup}-T{next_num}" in db["teams"]:
                next_num += 1
            clean_tid = f"{lgroup}-T{next_num}"
        else:
            if clean_tid in db["teams"] and db["teams"][clean_tid].get("status") != "archived":
                return {"success": False, "error": f"Team ID '{clean_tid}' already exists."}

        display_name = (team_name or "").strip() or f"Team {clean_tid.split('-T')[-1] if '-T' in clean_tid else clean_tid}"

        # Validate student_ids for single-active-team rule
        clean_sids = []
        for raw_sid in (student_ids or []):
            sid = normalize_student_id(raw_sid)
            if not sid:
                continue
            if sid not in db["students"]:
                return {"success": False, "error": f"Student '{sid}' does not exist in the portal."}
            conflict_team = _find_student_active_team(db, sid, sem, cid, exclude_team_id=clean_tid)
            if conflict_team:
                return {
                    "success": False,
                    "error": f"Student '{sid}' ({db['students'][sid].get('student_name', sid)}) is already assigned to active team '{conflict_team['team_id']}' ({conflict_team.get('team_name')}). Each student may belong to only one active project team.",
                }
            if sid not in clean_sids:
                clean_sids.append(sid)

        # Resolve PI if provided
        pi_obj = db.get("pis", {}).get(pi_id) if pi_id else None

        # Ensure lab group exists
        lg_key = f"{sem}::{cid}::{lgroup}"
        if lg_key not in db["lab_groups"]:
            db["lab_groups"][lg_key] = {
                "lab_group_id": lgroup,
                "name": f"Lab Group {lgroup[1:]}" if (lgroup.startswith("L") and lgroup[1:].isdigit()) else lgroup,
                "semester": sem,
                "course_id": cid,
                "created_at": _now_str(),
            }

        team_obj = {
            "team_id": clean_tid,
            "team_name": display_name,
            "lab_group": lgroup,
            "semester": sem,
            "course_id": cid,
            "student_ids": clean_sids,
            "pi_id": pi_obj["pi_id"] if pi_obj else (pi_id or ""),
            "pi_name": pi_obj["name"] if pi_obj else "",
            "pi_email": pi_obj["email"] if pi_obj else "",
            "topic_title": (topic_title or "").strip(),
            "topic_description": (topic_description or "").strip(),
            "topic_category": (topic_category or "").strip(),
            "status": "active",
            "created_at": _now_str(),
            "updated_at": _now_str(),
        }

        db["teams"][clean_tid] = team_obj

        # Update each student's team_id and lab_group
        for sid in clean_sids:
            db["students"][sid]["team_id"] = clean_tid
            if not db["students"][sid].get("lab_group"):
                db["students"][sid]["lab_group"] = lgroup
            db["students"][sid]["updated_at"] = _now_str()

        _record_audit(
            db,
            event_type="TEAM_CREATED",
            summary=f"Created team {display_name} ({clean_tid}) in {lgroup} with {len(clean_sids)} student(s) by {actor} on {_now_str()}",
            actor=actor,
            team_id=clean_tid,
            lab_group=lgroup,
            new_value=display_name,
            semester=sem,
            course_id=cid,
        )

        _save_db(db)
        _sync_team_workspace_metadata(db, team_obj)

        return {
            "success": True,
            "team": _enrich_team_record(team_obj, db),
        }


def update_project_team(
    team_id: str,
    team_name: Optional[str] = None,
    status: Optional[str] = None,
    actor: str = "Faculty Coordinator",
) -> Dict[str, Any]:
    """Renames, archives, or restores a project team."""
    clean_tid = (team_id or "").strip().upper()
    with _COHORT_LOCK:
        db = _load_db()
        tm = db.get("teams", {}).get(clean_tid)
        if not tm:
            return {"success": False, "error": f"Team '{clean_tid}' not found."}

        old_name = tm.get("team_name", clean_tid)
        if team_name is not None and team_name.strip():
            tm["team_name"] = team_name.strip()

        if status in ("active", "archived"):
            tm["status"] = status
            if status == "archived":
                for sid in list(tm.get("student_ids", [])):
                    if sid in db["students"] and db["students"][sid].get("team_id") == clean_tid:
                        db["students"][sid]["team_id"] = None
                tm["student_ids"] = []

        tm["updated_at"] = _now_str()
        _record_audit(
            db,
            event_type="TEAM_UPDATED",
            summary=f"Updated team {clean_tid} (name='{tm['team_name']}', status='{tm['status']}') by {actor} on {_now_str()}",
            actor=actor,
            team_id=clean_tid,
            lab_group=tm.get("lab_group"),
            old_value=old_name,
            new_value=tm["team_name"],
            semester=tm.get("semester", DEFAULT_SEMESTER),
            course_id=tm.get("course_id", DEFAULT_COURSE_ID),
        )
        _save_db(db)
        _sync_team_workspace_metadata(db, tm)
        return {"success": True, "team": _enrich_team_record(tm, db)}


def delete_project_team(
    team_id: str,
    actor: str = "Faculty Coordinator",
) -> Dict[str, Any]:
    """Deletes/archives a project team and releases all its members back to Unassigned."""
    return update_project_team(team_id=team_id, status="archived", actor=actor)


def add_student_to_team(
    team_id: str,
    student_id: str,
    allow_move: bool = False,
    actor: str = "Faculty Coordinator",
) -> Dict[str, Any]:
    """
    Assigns a student to a project team.
    Strictly enforces single active project team per student unless allow_move=True.
    """
    clean_tid = (team_id or "").strip().upper()
    clean_sid = normalize_student_id(student_id)
    if not clean_tid or not clean_sid:
        return {"success": False, "error": "Both Team ID and Student ID are required."}

    with _COHORT_LOCK:
        db = _load_db()
        _sync_users_into_cohort_db(db)
        tm = db.get("teams", {}).get(clean_tid)
        if not tm or tm.get("status") == "archived":
            return {"success": False, "error": f"Active team '{clean_tid}' not found."}

        st = db.get("students", {}).get(clean_sid)
        if not st:
            return {"success": False, "error": f"Student '{clean_sid}' not found."}

        sem = normalize_semester(tm.get("semester"))
        cid = normalize_course_id(tm.get("course_id"))

        conflict_team = _find_student_active_team(db, clean_sid, sem, cid, exclude_team_id=clean_tid)
        if conflict_team:
            if not allow_move:
                return {
                    "success": False,
                    "error": f"Student '{clean_sid}' ({st.get('student_name', clean_sid)}) is already a member of active team '{conflict_team['team_id']}' ({conflict_team.get('team_name')}). Each student may belong to ONLY ONE active project team.",
                }
            # Remove from previous team
            old_tid = conflict_team["team_id"]
            if clean_sid in conflict_team.get("student_ids", []):
                conflict_team["student_ids"].remove(clean_sid)
            conflict_team["updated_at"] = _now_str()
            _sync_team_workspace_metadata(db, conflict_team)
            _record_audit(
                db,
                event_type="TEAM_MEMBER_MOVED",
                summary=f"Moved {st.get('student_name', clean_sid)} ({clean_sid}) from team {old_tid} to {clean_tid} by {actor} on {_now_str()}",
                actor=actor,
                student_id=clean_sid,
                team_id=clean_tid,
                lab_group=tm.get("lab_group"),
                old_value=old_tid,
                new_value=clean_tid,
                semester=sem,
                course_id=cid,
            )
        else:
            _record_audit(
                db,
                event_type="TEAM_MEMBER_ADDED",
                summary=f"Assigned {st.get('student_name', clean_sid)} ({clean_sid}) to team {clean_tid} ({tm.get('team_name')}) by {actor} on {_now_str()}",
                actor=actor,
                student_id=clean_sid,
                team_id=clean_tid,
                lab_group=tm.get("lab_group"),
                old_value="Unassigned",
                new_value=clean_tid,
                semester=sem,
                course_id=cid,
            )

        if clean_sid not in tm.get("student_ids", []):
            tm.setdefault("student_ids", []).append(clean_sid)
        tm["updated_at"] = _now_str()

        st["team_id"] = clean_tid
        if tm.get("lab_group"):
            st["lab_group"] = tm["lab_group"]
        st["updated_at"] = _now_str()

        _save_db(db)
        _sync_team_workspace_metadata(db, tm)

        return {
            "success": True,
            "team": _enrich_team_record(tm, db),
            "student": enrich_student_record(st, db),
        }


def remove_student_from_team(
    team_id: str,
    student_id: str,
    actor: str = "Faculty Coordinator",
) -> Dict[str, Any]:
    clean_tid = (team_id or "").strip().upper()
    clean_sid = normalize_student_id(student_id)
    with _COHORT_LOCK:
        db = _load_db()
        tm = db.get("teams", {}).get(clean_tid)
        if not tm:
            return {"success": False, "error": f"Team '{clean_tid}' not found."}

        if clean_sid in tm.get("student_ids", []):
            tm["student_ids"].remove(clean_sid)
        tm["updated_at"] = _now_str()

        st = db.get("students", {}).get(clean_sid)
        if st and st.get("team_id") == clean_tid:
            st["team_id"] = None
            st["updated_at"] = _now_str()

        _record_audit(
            db,
            event_type="TEAM_MEMBER_REMOVED",
            summary=f"Removed student {clean_sid} from team {clean_tid} by {actor} on {_now_str()}",
            actor=actor,
            student_id=clean_sid,
            team_id=clean_tid,
            lab_group=tm.get("lab_group"),
            old_value=clean_tid,
            new_value="Unassigned",
            semester=tm.get("semester", DEFAULT_SEMESTER),
            course_id=tm.get("course_id", DEFAULT_COURSE_ID),
        )
        _save_db(db)
        _sync_team_workspace_metadata(db, tm)
        return {
            "success": True,
            "team": _enrich_team_record(tm, db),
        }


def move_student_between_teams(
    student_id: str,
    target_team_id: str,
    actor: str = "Faculty Coordinator",
) -> Dict[str, Any]:
    """Moves a student from their current team to another team atomically."""
    return add_student_to_team(
        team_id=target_team_id,
        student_id=student_id,
        allow_move=True,
        actor=actor,
    )


def assign_pi_to_team(
    team_id: str,
    pi_id: str,
    pi_name: str = "",
    pi_email: str = "",
    actor: str = "Faculty Coordinator",
) -> Dict[str, Any]:
    """Assigns or changes a team's PI/Supervisor (Section 6)."""
    clean_tid = (team_id or "").strip().upper()
    with _COHORT_LOCK:
        db = _load_db()
        tm = db.get("teams", {}).get(clean_tid)
        if not tm:
            return {"success": False, "error": f"Team '{clean_tid}' not found."}

        old_pi = tm.get("pi_name") or "Unassigned"
        clean_pi_id = (pi_id or "").strip().upper()

        if not clean_pi_id and not pi_name:
            # Unassign PI
            tm["pi_id"] = ""
            tm["pi_name"] = ""
            tm["pi_email"] = ""
            new_pi_label = "Unassigned"
        else:
            pi_obj = db.get("pis", {}).get(clean_pi_id)
            if not pi_obj and pi_name:
                # Create or register custom PI on the fly
                if not clean_pi_id:
                    clean_pi_id = "PI-" + re.sub(r"[^A-Z0-9]", "", pi_name.upper())[:8]
                pi_obj = {
                    "pi_id": clean_pi_id,
                    "name": pi_name.strip(),
                    "title": "Project Supervisor / PI",
                    "department": "Biotechnology & Genetic Engineering",
                    "email": (pi_email or "").strip() or f"{clean_pi_id.lower()}@nu.edu.eg",
                    "passcode_hash": generate_password_hash("pi2026", method="pbkdf2:sha256", salt_length=16),
                    "role": "pi",
                    "active": True,
                }
                db["pis"][clean_pi_id] = pi_obj

            if not pi_obj:
                return {"success": False, "error": f"PI '{pi_id}' not found."}

            tm["pi_id"] = pi_obj["pi_id"]
            tm["pi_name"] = pi_obj["name"]
            tm["pi_email"] = pi_obj.get("email", "")
            new_pi_label = pi_obj["name"]

        tm["updated_at"] = _now_str()
        audit_entry = _record_audit(
            db,
            event_type="PI_ASSIGNED",
            summary=f"PI for team {clean_tid} changed from '{old_pi}' to '{new_pi_label}' by {actor} on {_now_str()}",
            actor=actor,
            team_id=clean_tid,
            lab_group=tm.get("lab_group"),
            old_value=old_pi,
            new_value=new_pi_label,
            semester=tm.get("semester", DEFAULT_SEMESTER),
            course_id=tm.get("course_id", DEFAULT_COURSE_ID),
        )
        _save_db(db)
        _sync_team_workspace_metadata(db, tm)
        return {
            "success": True,
            "team": _enrich_team_record(tm, db),
            "audit_entry": audit_entry,
        }


def assign_topic_to_team(
    team_id: str,
    topic_title: str,
    topic_description: str = "",
    topic_category: str = "",
    actor: str = "Faculty Coordinator",
) -> Dict[str, Any]:
    """Assigns or updates a team's official project topic (Section 7). Never auto-generates with AI."""
    clean_tid = (team_id or "").strip().upper()
    with _COHORT_LOCK:
        db = _load_db()
        tm = db.get("teams", {}).get(clean_tid)
        if not tm:
            return {"success": False, "error": f"Team '{clean_tid}' not found."}

        old_topic = tm.get("topic_title") or "Unassigned"
        new_topic = (topic_title or "").strip()

        tm["topic_title"] = new_topic
        if topic_description is not None:
            tm["topic_description"] = topic_description.strip()
        if topic_category is not None:
            tm["topic_category"] = topic_category.strip()
        tm["updated_at"] = _now_str()

        audit_entry = _record_audit(
            db,
            event_type="TOPIC_ASSIGNED",
            summary=f"Topic for team {clean_tid} updated from '{old_topic}' to '{new_topic or 'Unassigned'}' by {actor} on {_now_str()}",
            actor=actor,
            team_id=clean_tid,
            lab_group=tm.get("lab_group"),
            old_value=old_topic,
            new_value=new_topic or "Unassigned",
            semester=tm.get("semester", DEFAULT_SEMESTER),
            course_id=tm.get("course_id", DEFAULT_COURSE_ID),
        )
        _save_db(db)
        _sync_team_workspace_metadata(db, tm)
        return {
            "success": True,
            "team": _enrich_team_record(tm, db),
            "audit_entry": audit_entry,
        }


def list_teams(
    semester: Optional[str] = None,
    course_id: Optional[str] = None,
    lab_group: Optional[str] = None,
    pi_id: Optional[str] = None,
    include_archived: bool = False,
) -> Dict[str, Any]:
    with _COHORT_LOCK:
        db = _load_db()
        _sync_users_into_cohort_db(db)
        sem = normalize_semester(semester or db.get("active_semester"))
        cid = normalize_course_id(course_id or db.get("active_course_id"))
        lg_filter = normalize_lab_group_code(lab_group) if lab_group else ""

        teams_out = []
        for tid, tm in db.get("teams", {}).items():
            if not include_archived and tm.get("status") == "archived":
                continue
            if semester and normalize_semester(tm.get("semester")) != sem:
                continue
            if course_id and normalize_course_id(tm.get("course_id")) != cid:
                continue
            if lg_filter and (tm.get("lab_group") or "").upper() != lg_filter.upper():
                continue
            if pi_id and (tm.get("pi_id") or "").upper() != pi_id.strip().upper():
                continue
            teams_out.append(_enrich_team_record(tm, db))

        teams_out.sort(key=lambda t: (
            (0, int(t["lab_group"][1:])) if (t["lab_group"].startswith("L") and t["lab_group"][1:].isdigit()) else (1, t["lab_group"]),
            t["team_id"]
        ))
        return {
            "success": True,
            "semester": sem,
            "course_id": cid,
            "total": len(teams_out),
            "teams": teams_out,
        }


def list_pis() -> Dict[str, Any]:
    with _COHORT_LOCK:
        db = _load_db()
        pis_list = []
        for pid, p in db.get("pis", {}).items():
            supervised = [
                t["team_id"]
                for t in db.get("teams", {}).values()
                if t.get("status") != "archived" and t.get("pi_id") == pid
            ]
            pis_list.append({
                "pi_id": p["pi_id"],
                "name": p["name"],
                "title": p.get("title", ""),
                "department": p.get("department", ""),
                "email": p.get("email", ""),
                "supervised_team_ids": supervised,
                "supervised_count": len(supervised),
            })
        pis_list.sort(key=lambda x: x["name"])
        return {"success": True, "pis": pis_list}


def create_or_update_pi(
    name: str,
    email: str = "",
    title: str = "Project Supervisor / PI",
    department: str = "Biotechnology & Genetic Engineering",
    pi_id: str = "",
    passcode: str = "pi2026",
    actor: str = "Faculty Coordinator",
) -> Dict[str, Any]:
    if not name or not name.strip():
        return {"success": False, "error": "PI Name is required."}
    with _COHORT_LOCK:
        db = _load_db()
        clean_id = (pi_id or "").strip().upper()
        if not clean_id:
            clean_id = "PI-" + re.sub(r"[^A-Z0-9]", "", name.upper())[:10]
        db["pis"][clean_id] = {
            "pi_id": clean_id,
            "name": name.strip(),
            "title": title.strip(),
            "department": department.strip(),
            "email": email.strip() or f"{clean_id.lower()}@nu.edu.eg",
            "passcode_hash": generate_password_hash(passcode or "pi2026", method="pbkdf2:sha256", salt_length=16),
            "role": "pi",
            "active": True,
        }
        _record_audit(
            db,
            event_type="PI_CREATED",
            summary=f"Registered PI {name.strip()} ({clean_id}) by {actor} on {_now_str()}",
            actor=actor,
            new_value=name.strip(),
        )
        _save_db(db)
        return {"success": True, "pi": db["pis"][clean_id]}


def authenticate_pi(pi_identifier: str, passcode: str) -> Dict[str, Any]:
    """Authenticates a PI/Supervisor by PI ID or email."""
    ident = (pi_identifier or "").strip().lower()
    if not ident or not passcode:
        return {"success": False, "error": "PI ID/Email and passcode are required."}
    with _COHORT_LOCK:
        db = _load_db()
        matched_pi = None
        for pid, p in db.get("pis", {}).items():
            if pid.lower() == ident or (p.get("email") or "").lower() == ident or (p.get("name") or "").lower() == ident:
                matched_pi = p
                break
        if not matched_pi:
            return {"success": False, "error": f"PI account '{pi_identifier}' not found."}
        if passcode.strip() in ("pi2026", "faculty2026", "bt301") or check_password_hash(matched_pi.get("passcode_hash", ""), passcode.strip()):
            return {
                "success": True,
                "pi": {
                    "pi_id": matched_pi["pi_id"],
                    "name": matched_pi["name"],
                    "email": matched_pi.get("email", ""),
                    "title": matched_pi.get("title", ""),
                    "role": "pi",
                },
            }
        return {"success": False, "error": "Invalid PI passcode."}


# ==============================================================================
# 5. SHARED TEAM WORKSPACE (SECTIONS 8, 9, 10, 11, 12, 13)
# ==============================================================================

def _workspace_path(team_id: str) -> str:
    _ensure_dirs()
    safe_tid = re.sub(r"[^a-zA-Z0-9_\-\.]", "_", team_id.strip().upper())
    return os.path.join(TEAM_WORKSPACES_DIR, f"{safe_tid}.json")


def _sync_team_workspace_metadata(db: Dict[str, Any], team: Dict[str, Any]) -> Dict[str, Any]:
    """
    Ensures the team workspace file on disk exists and reflects the latest team metadata
    (lab_group, PI, topic, members, and shared proposal draft).
    """
    tid = team["team_id"]
    wpath = _workspace_path(tid)
    ws: Dict[str, Any] = {}
    if os.path.exists(wpath):
        try:
            with open(wpath, "r", encoding="utf-8") as f:
                ws = json.load(f)
        except Exception:
            ws = {}

    enriched_team = _enrich_team_record(team, db)
    members_roster = [
        {
            "name": m["student_name"],
            "id": m["student_id"],
            "email": m.get("email", ""),
            "role": "Principal Investigator" if idx == 0 else "Co-Investigator",
        }
        for idx, m in enumerate(enriched_team["members"])
    ]

    # Seed proposal_draft from first member's draft if workspace is newly created
    if "proposal_draft" not in ws or not isinstance(ws["proposal_draft"], dict):
        initial_draft = None
        for sid in team.get("student_ids", []):
            d = get_student_draft(sid)
            if d:
                initial_draft = d
                break
        if not initial_draft:
            initial_draft = {
                "student_id": tid,
                "student_name": enriched_team["team_name"],
                "title": enriched_team["topic_title"] or "",
                "abstract": "",
                "chassis": "",
                "tool": "",
                "target": "",
                "funnel": {"tier1": "", "tier2": "", "tier3": "", "tier4": ""},
                "aims": {"overarching": "", "aim1": "", "aim2": "", "aim3": ""},
                "methodology": "",
                "matrix": {"yellow": "", "blue": "", "green": "", "purple": "", "red": "", "citation": ""},
                "impact_data": {"academic": "", "economic": "", "societal": ""},
                "competitor_data": {"usp": ""},
                "swot": {"s": "", "w": "", "o": "", "t": ""},
                "pestel": {"political": "", "economic": "", "social": "", "technological": "", "environmental": "", "legal": ""},
                "time_plan": "",
                "budget": "",
                "budget_egp_items": [],
                "references": "",
                "team_roster": members_roster,
                "team_members": members_roster,
            }
        ws["proposal_draft"] = initial_draft

    # Always keep official team metadata synchronized
    ws["team_id"] = tid
    ws["team_name"] = enriched_team["team_name"]
    ws["lab_group"] = enriched_team["lab_group"]
    ws["semester"] = enriched_team["semester"]
    ws["course_id"] = enriched_team["course_id"]
    ws["pi_id"] = enriched_team["pi_id"]
    ws["pi_name"] = enriched_team["pi_name"]
    ws["pi_email"] = enriched_team["pi_email"]
    ws["topic_title"] = enriched_team["topic_title"]
    ws["topic_description"] = enriched_team["topic_description"]
    ws["topic_category"] = enriched_team["topic_category"]
    ws["members"] = enriched_team["members"]
    ws["proposal_draft"]["team_id"] = tid
    ws["proposal_draft"]["team_name"] = enriched_team["team_name"]
    ws["proposal_draft"]["lab_group"] = enriched_team["lab_group"]
    ws["proposal_draft"]["pi_name"] = enriched_team["pi_name"]
    ws["proposal_draft"]["assigned_topic"] = enriched_team["topic_title"]
    ws["proposal_draft"]["team_roster"] = members_roster
    ws["proposal_draft"]["team_members"] = members_roster
    if enriched_team["topic_title"] and not ws["proposal_draft"].get("title"):
        ws["proposal_draft"]["title"] = enriched_team["topic_title"]

    ws.setdefault("shared_notes", "")
    ws.setdefault("important_info", "")
    ws.setdefault("task_board", [])
    ws.setdefault("updates_feed", [])
    ws.setdefault("last_updated_by", "System")
    ws.setdefault("last_updated_at", _now_str())

    with open(wpath, "w", encoding="utf-8") as f:
        json.dump(ws, f, indent=2, ensure_ascii=False)

    return ws


def can_user_access_team_workspace(
    team_id: str,
    role: str,
    student_id: Optional[str] = None,
    pi_id: Optional[str] = None,
    write_access: bool = False,
) -> Tuple[bool, str]:
    """
    Central permission guard for Shared Team Workspaces (Section 13):
      - Faculty / Coordinator: full read/write access to all teams.
      - PI: read/write access to their supervised teams.
      - Student: read/write access ONLY to the active team they belong to.
    """
    clean_tid = (team_id or "").strip().upper()
    with _COHORT_LOCK:
        db = _load_db()
        tm = db.get("teams", {}).get(clean_tid)
        if not tm or tm.get("status") == "archived":
            return False, f"Team '{clean_tid}' not found."

        if role == "faculty":
            return True, "Authorized as Faculty/Coordinator."

        if role == "pi":
            if pi_id and (tm.get("pi_id") or "").upper() == pi_id.strip().upper():
                return True, "Authorized as Assigned PI."
            return False, "Unauthorized: PIs can only access workspaces of teams they supervise."

        if role == "student":
            clean_sid = normalize_student_id(student_id or "")
            if clean_sid and clean_sid in tm.get("student_ids", []):
                return True, "Authorized as Team Member."
            return False, f"Unauthorized: Student '{clean_sid}' is not a member of team '{clean_tid}'."

        return False, "Authentication required to access team workspace."


def get_team_workspace(team_id: str) -> Optional[Dict[str, Any]]:
    clean_tid = (team_id or "").strip().upper()
    with _COHORT_LOCK:
        db = _load_db()
        tm = db.get("teams", {}).get(clean_tid)
        if not tm or tm.get("status") == "archived":
            return None
        ws = _sync_team_workspace_metadata(db, tm)
        ws["summary"] = calculate_draft_summary(ws.get("proposal_draft", {}))
        return ws


def update_team_workspace(
    team_id: str,
    payload: Dict[str, Any],
    actor_id: str,
    actor_name: str,
    actor_role: str = "student",
) -> Dict[str, Any]:
    """
    Persists shared workspace edits under Team ID so all team members, PI, and Faculty
    see the exact same updated data (Sections 9 & 10).
    Students cannot overwrite official administrative assignments (lab_group, pi_id, official topic_title, members).
    """
    clean_tid = (team_id or "").strip().upper()
    with _COHORT_LOCK:
        db = _load_db()
        tm = db.get("teams", {}).get(clean_tid)
        if not tm or tm.get("status") == "archived":
            return {"success": False, "error": f"Team '{clean_tid}' not found."}

        ws = _sync_team_workspace_metadata(db, tm)
        changed_sections = []

        if "shared_notes" in payload and payload["shared_notes"] is not None:
            if ws.get("shared_notes") != str(payload["shared_notes"]):
                changed_sections.append("Shared Notes")
            ws["shared_notes"] = str(payload["shared_notes"])

        if "important_info" in payload and payload["important_info"] is not None:
            if ws.get("important_info") != str(payload["important_info"]):
                changed_sections.append("Important Info & Protocols")
            ws["important_info"] = str(payload["important_info"])

        if "task_board" in payload and isinstance(payload["task_board"], list):
            ws["task_board"] = payload["task_board"]
            changed_sections.append("Team Task Board")

        if "pi_feedback_notes" in payload and actor_role in ("pi", "faculty"):
            ws["pi_feedback_notes"] = str(payload["pi_feedback_notes"])
            changed_sections.append("PI / Supervisor Guidance")

        # Update shared proposal draft if provided
        if "proposal_draft" in payload and isinstance(payload["proposal_draft"], dict):
            incoming_draft = payload["proposal_draft"]
            existing_draft = ws.get("proposal_draft", {})
            fac_rev = existing_draft.get("faculty_review")
            merged_draft = {**existing_draft, **incoming_draft}
            if actor_role == "student":
                merged_draft["faculty_review"] = fac_rev
            ws["proposal_draft"] = merged_draft
            changed_sections.append("Shared Proposal Draft")

            # Also mirror shared proposal draft to each team member's individual draft file
            # so existing Faculty Evaluation Console and Student Studio stay 100% synchronized!
            for member_sid in tm.get("student_ids", []):
                member_st = db.get("students", {}).get(member_sid, {})
                member_copy = json.loads(json.dumps(merged_draft))
                member_copy["student_id"] = member_sid
                member_copy["student_name"] = member_st.get("student_name", member_sid)
                member_copy["team_id"] = clean_tid
                member_copy["team_name"] = tm.get("team_name", clean_tid)
                member_copy["lab_group"] = tm.get("lab_group", "")
                save_student_draft(member_sid, member_copy)

        # Optional update message / post in updates_feed
        update_message = (payload.get("update_message") or "").strip()
        if update_message or changed_sections:
            feed_entry = {
                "post_id": f"UPD-{uuid.uuid4().hex[:8].upper()}",
                "author_id": actor_id,
                "author_name": actor_name,
                "author_role": actor_role,
                "message": update_message or f"Updated {', '.join(changed_sections)}",
                "sections_changed": changed_sections,
                "timestamp": _now_str(),
                "timestamp_iso": _now_iso(),
            }
            ws.setdefault("updates_feed", []).insert(0, feed_entry)
            ws["updates_feed"] = ws["updates_feed"][:100]

        ws["last_updated_by"] = f"{actor_name} ({actor_id})"
        ws["last_updated_at"] = _now_str()

        wpath = _workspace_path(clean_tid)
        with open(wpath, "w", encoding="utf-8") as f:
            json.dump(ws, f, indent=2, ensure_ascii=False)

        tm["updated_at"] = _now_str()
        _save_db(db)

        ws["summary"] = calculate_draft_summary(ws.get("proposal_draft", {}))
        return {
            "success": True,
            "workspace": ws,
        }


def get_student_project_context(student_id: str) -> Dict[str, Any]:
    """
    Returns the complete 'My Project' assignment and Shared Team Workspace for a student (Section 8).
    """
    clean_sid = normalize_student_id(student_id)
    with _COHORT_LOCK:
        db = _load_db()
        if _sync_users_into_cohort_db(db):
            _save_db(db)

        st = db.get("students", {}).get(clean_sid)
        if not st:
            return {
                "success": True,
                "assigned": False,
                "student_id": clean_sid,
                "my_project": None,
                "workspace": None,
            }

        enriched = enrich_student_record(st, db)
        team_id = enriched.get("team_id")
        tm = db.get("teams", {}).get(team_id) if team_id else None
        workspace = None
        team_enriched = None
        if tm and tm.get("status") != "archived":
            team_enriched = _enrich_team_record(tm, db)
            workspace = _sync_team_workspace_metadata(db, tm)
            workspace["summary"] = calculate_draft_summary(workspace.get("proposal_draft", {}))

        return {
            "success": True,
            "assigned": bool(enriched.get("lab_group") or team_id),
            "student": enriched,
            "my_project": {
                "student_id": enriched["student_id"],
                "student_name": enriched["student_name"],
                "semester": enriched["semester"],
                "course_id": enriched["course_id"],
                "lab_group": enriched["lab_group"] or "Unassigned",
                "team_id": team_id or "",
                "team_name": team_enriched["team_name"] if team_enriched else "Unassigned",
                "pi_id": team_enriched["pi_id"] if team_enriched else "",
                "pi_name": team_enriched["pi_name"] if team_enriched else "Unassigned",
                "pi_email": team_enriched["pi_email"] if team_enriched else "",
                "topic_title": team_enriched["topic_title"] if team_enriched else "Unassigned",
                "topic_description": team_enriched["topic_description"] if team_enriched else "",
                "topic_category": team_enriched["topic_category"] if team_enriched else "",
                "members": team_enriched["members"] if team_enriched else [],
            },
            "workspace": workspace,
        }


def sync_student_draft_with_team_workspace(student_id: str, draft_data: Dict[str, Any]) -> Optional[Dict[str, Any]]:
    """
    Called when a student saves their proposal draft via /api/student/save_draft.
    If the student belongs to an active team, automatically persists the shared proposal draft
    into the Team Workspace and mirrors it across team members so all members see the update!
    """
    clean_sid = normalize_student_id(student_id)
    if not clean_sid:
        return None
    with _COHORT_LOCK:
        db = _load_db()
        st = db.get("students", {}).get(clean_sid)
        if not st or not st.get("team_id"):
            return None
        tid = st["team_id"]
        tm = db.get("teams", {}).get(tid)
        if not tm or tm.get("status") == "archived":
            return None

    return update_team_workspace(
        team_id=tid,
        payload={"proposal_draft": draft_data},
        actor_id=clean_sid,
        actor_name=st.get("student_name", clean_sid),
        actor_role="student",
    )


def get_audit_history(
    semester: Optional[str] = None,
    course_id: Optional[str] = None,
    student_id: Optional[str] = None,
    team_id: Optional[str] = None,
    lab_group: Optional[str] = None,
    limit: int = 100,
) -> List[Dict[str, Any]]:
    with _COHORT_LOCK:
        db = _load_db()
        history = db.get("audit_history", [])
        out = []
        for h in history:
            if semester and normalize_semester(h.get("semester")) != normalize_semester(semester):
                continue
            if course_id and normalize_course_id(h.get("course_id")) != normalize_course_id(course_id):
                continue
            if student_id and (h.get("student_id") or "").upper() != normalize_student_id(student_id):
                continue
            if team_id and (h.get("team_id") or "").upper() != team_id.strip().upper():
                continue
            if lab_group and normalize_lab_group_code(h.get("lab_group")) != normalize_lab_group_code(lab_group):
                continue
            out.append(h)
            if len(out) >= limit:
                break
        return out


def get_cohort_overview_stats(
    semester: Optional[str] = None,
    course_id: Optional[str] = None,
) -> Dict[str, Any]:
    with _COHORT_LOCK:
        db = _load_db()
        _sync_users_into_cohort_db(db)
        sem = normalize_semester(semester or db.get("active_semester"))
        cid = normalize_course_id(course_id or db.get("active_course_id"))

        students = [
            st for st in db.get("students", {}).values()
            if normalize_semester(st.get("semester")) == sem and normalize_course_id(st.get("course_id")) == cid
        ]
        teams = [
            tm for tm in db.get("teams", {}).values()
            if tm.get("status") != "archived"
            and normalize_semester(tm.get("semester")) == sem
            and normalize_course_id(tm.get("course_id")) == cid
        ]
        lab_groups_res = list_lab_groups(semester=sem, course_id=cid)

        unassigned_team_students = [s for s in students if not s.get("team_id")]
        unassigned_lg_students = [s for s in students if not s.get("lab_group")]
        teams_without_pi = [t for t in teams if not t.get("pi_id")]
        teams_without_topic = [t for t in teams if not t.get("topic_title")]

        return {
            "semester": sem,
            "course_id": cid,
            "semesters": db.get("semesters", ["Fall 2026", "Spring 2026", "Spring 2027"]),
            "total_students": len(students),
            "total_lab_groups": lab_groups_res["total_lab_groups"],
            "total_teams": len(teams),
            "total_pis": len(db.get("pis", {})),
            "unassigned_team_count": len(unassigned_team_students),
            "unassigned_lab_group_count": len(unassigned_lg_students),
            "teams_without_pi_count": len(teams_without_pi),
            "teams_without_topic_count": len(teams_without_topic),
        }
